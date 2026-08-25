from __future__ import annotations

import importlib.util
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
BASE_PATH = ROOT / ".codex-tmp" / "create_cloudflare_vs_selfhost_report.py"
OUTPUT_DIR = ROOT / "outputs"
OUTPUT_PATH = OUTPUT_DIR / "bao-cao-cloudflare-vs-vertex-ai-bang-chi-phi-moi.docx"

spec = importlib.util.spec_from_file_location("base_report", BASE_PATH)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)


USD_VND = 26_000
DAYS = 30
MINUTES_PER_USER_DAY = 10
MINUTES_PER_USER_MONTH = DAYS * MINUTES_PER_USER_DAY
HOURS_PER_USER_MONTH = MINUTES_PER_USER_MONTH / 60

CF_NEURONS_PER_MINUTE = 46.63
CF_FREE_NEURONS_DAY = 10_000
CF_PRICE_PER_1000_NEURONS = 0.011
CF_PAID_BASE = 5.0
CF_NEURONS_USER_DAY = MINUTES_PER_USER_DAY * CF_NEURONS_PER_MINUTE
CF_FREE_USERS = int(CF_FREE_NEURONS_DAY // CF_NEURONS_USER_DAY)

VERTEX_CPU_HOUR = 0.219
VERTEX_T4_HOUR = 0.402
VERTEX_NODE_HOUR = VERTEX_CPU_HOUR + VERTEX_T4_HOUR
VERTEX_WARM_MONTH = VERTEX_NODE_HOUR * 730
VERTEX_HA_MONTH = VERTEX_WARM_MONTH * 2
VERTEX_DAILY_FLOOR = VERTEX_NODE_HOUR * DAYS

SCALES = [1, 100, 500, 1000]


def money(value: float) -> str:
    return f"${value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def vnd_short(value_usd: float) -> str:
    v = value_usd * USD_VND
    if v >= 1_000_000:
        return f"{v / 1_000_000:.2f} triệu".replace(".", ",")
    if v >= 1_000:
        return f"{v / 1_000:.0f} nghìn"
    return f"{v:.0f}đ"


def cf_ai_cost(users: int) -> float:
    excess_day = max(0.0, CF_NEURONS_USER_DAY * users - CF_FREE_NEURONS_DAY)
    return excess_day * DAYS * CF_PRICE_PER_1000_NEURONS / 1000


def cf_total(users: int) -> float:
    if users <= CF_FREE_USERS:
        return 0.0
    return CF_PAID_BASE + cf_ai_cost(users)


def vertex_compute_range(users: int) -> tuple[float, float]:
    audio_hours = HOURS_PER_USER_MONTH * users
    return (
        audio_hours / 30 * VERTEX_NODE_HOUR,
        audio_hours / 10 * VERTEX_NODE_HOUR,
    )


def add_caption_with_links(doc, items: list[tuple[str, str]]) -> None:
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.space_before = base.Pt(4)
    p.paragraph_format.space_after = base.Pt(4)
    p.add_run("Nguồn: ")
    for idx, (label, url) in enumerate(items):
        if idx:
            p.add_run("; ")
        base.add_hyperlink(p, label, url)
    p.add_run(".")


def add_metadata(doc, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = base.Pt(3)
    base.set_run_font(p.add_run(f"{label}: "), size=10.5, color=base.NAVY, bold=True)
    base.set_run_font(p.add_run(value), size=10.5, color=base.INK)


def style_group_cell(cell, label: str, fill: str, accent: str) -> None:
    cell.text = label
    cell.vertical_alignment = base.WD_CELL_VERTICAL_ALIGNMENT.CENTER
    base.set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = base.Pt(0)
    p.paragraph_format.space_after = base.Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for run in p.runs:
        base.set_run_font(run, size=9.0, color=accent, bold=True)


def build_report() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    base.configure_styles(doc)
    base.configure_sections(doc)
    base.BULLET_NUM_ID, base.NUMBER_NUM_ID = base.add_custom_numbering(doc)

    props = doc.core_properties
    props.title = "Cloudflare Workers AI và Vertex AI - chi phí 10 phút mỗi ngày mỗi người dùng"
    props.subject = "So sánh chi phí ASR ở quy mô 1, 100, 500 và 1.000 người dùng"
    props.author = "Speaking AI"
    props.last_modified_by = "Speaking AI"
    props.keywords = "Cloudflare Workers AI, Vertex AI, Whisper, ASR, T4, chi phí"
    props.created = datetime(2026, 8, 21, tzinfo=timezone.utc)

    # Trang 1 - executive decision memo.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = base.Pt(18)
    p.paragraph_format.space_after = base.Pt(4)
    base.set_run_font(p.add_run("BÁO CÁO QUYẾT ĐỊNH HẠ TẦNG ASR"), size=10.5, color=base.BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = base.Pt(6)
    base.set_run_font(p.add_run("Cloudflare hay Vertex AI?"), size=25, color=base.NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = base.Pt(16)
    base.set_run_font(p.add_run("Chi phí khi mỗi người dùng nói 10 phút mỗi ngày"), size=14, color=base.MUTED)

    add_metadata(doc, "Khối lượng mỗi user", "10 phút/ngày x 30 ngày = 300 phút/tháng = 5 giờ audio")
    add_metadata(doc, "Cloudflare", "Workers AI + @cf/openai/whisper-large-v3-turbo")
    add_metadata(doc, "Vertex", "Custom endpoint: n1-standard-4 + NVIDIA Tesla T4 tại us-central1")
    add_metadata(doc, "Ngày giá", "21/08/2026; quy đổi 26.000 VND/USD; chưa gồm thuế")

    base.add_callout(
        doc,
        "KẾT LUẬN",
        "Từ 1 đến 1.000 user, Cloudflare là lựa chọn mặc định: rẻ hơn, triển khai nhanh và không phải vận hành GPU. Chỉ chuyển sang Vertex AI khi model tùy biến tạo ra lợi ích chất lượng rõ rệt, có yêu cầu IAM/VPC của GCP, hoặc benchmark thực tế chứng minh tổng chi phí thấp hơn.",
        fill=base.PALE_GREEN,
        accent=base.GREEN,
    )

    base.add_heading(doc, "Chi phí theo quy mô", 1)
    rows = []
    for users in SCALES:
        rows.append([
            "Cloudflare Workers AI" if users == 1 else "",
            f"{users:,}".replace(",", "."),
            f"{MINUTES_PER_USER_MONTH * users:,}".replace(",", "."),
            money(cf_total(users)),
            vnd_short(cf_total(users)),
        ])
    for users in SCALES:
        rows.append([
            "Vertex AI - 1 replica warm" if users == 1 else "",
            f"{users:,}".replace(",", "."),
            f"{MINUTES_PER_USER_MONTH * users:,}".replace(",", "."),
            money(VERTEX_WARM_MONTH),
            vnd_short(VERTEX_WARM_MONTH),
        ])
    summary_table = base.add_table(
        doc,
        ["Giải pháp", "User", "Phút/tháng", "Tổng USD", "Quy đổi VND"],
        rows,
        [2400, 960, 1900, 1800, 2300],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        font_size=8.6,
        header_fill=base.LIGHT_BLUE,
    )
    cloudflare_cell = summary_table.cell(1, 0).merge(summary_table.cell(4, 0))
    vertex_cell = summary_table.cell(5, 0).merge(summary_table.cell(8, 0))
    style_group_cell(cloudflare_cell, "Cloudflare Workers AI", base.PALE_GREEN, base.GREEN)
    style_group_cell(vertex_cell, "Vertex AI\n1 replica warm", base.PALE_GOLD, base.GOLD)
    base.add_body(
        doc,
        "Cloudflare tương ứng khoảng $17,09 cho 100 user; $78,64 cho 500 user; và $155,58 cho 1.000 user. Vertex warm một replica là $453,33/tháng và cần benchmark để biết một T4 có đáp ứng được concurrency hay không.",
        italic=True,
        color=base.MUTED,
    )

    base.add_heading(doc, "Ngưỡng quyết định", 1)
    base.add_bullet(doc, f"Cloudflare Free đủ tối đa khoảng {CF_FREE_USERS} user ở đúng mức 10 phút/ngày, nếu không có tải Workers AI khác.", base.BULLET_NUM_ID, "Cloudflare Free:")
    base.add_bullet(doc, "Từ 22 user trở lên, cần Workers Paid để inference không bị dừng khi vượt quota Neurons hằng ngày.", base.BULLET_NUM_ID, "Workers Paid:")
    base.add_bullet(doc, "So với một Vertex replica chạy 24/7, điểm hòa vốn lý thuyết của Cloudflare vào khoảng 2.935 user; với hai replica HA là khoảng 5.881 user.", base.BULLET_NUM_ID, "Hòa vốn lý thuyết:")

    # Trang 2 - Cloudflare.
    base.add_page_break(doc)
    base.add_heading(doc, "1. Cloudflare Workers AI", 1)
    base.add_callout(
        doc,
        "CÔNG THỨC",
        "Mỗi user dùng 466,3 Neurons/ngày. Chi phí AI tháng = max(0; 466,3 x số user - 10.000) x 30 x 0,011 / 1.000. Khi vượt quota ngày, cộng tối thiểu 5 USD/tháng cho Workers Paid.",
        fill=base.LIGHT_GRAY,
        accent=base.BLUE,
    )

    cf_rows = []
    for users in SCALES:
        ai_cost = cf_ai_cost(users)
        total = cf_total(users)
        if users == 1:
            note = "Free: $0; Paid tùy chọn: $5"
        else:
            note = f"AI {money(ai_cost)} + Paid $5"
        cf_rows.append([
            f"{users:,}".replace(",", "."),
            f"{MINUTES_PER_USER_MONTH * users:,}".replace(",", "."),
            money(total),
            vnd_short(total),
            note,
        ])
    base.add_table(
        doc,
        ["User", "Phút/tháng", "Tổng USD", "Quy đổi VND", "Cấu phần"],
        cf_rows,
        [1000, 1700, 1600, 1900, 3160],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
        header_fill=base.PALE_GREEN,
    )
    add_caption_with_links(doc, [
        ("Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/"),
        ("Whisper Large V3 Turbo", "https://developers.cloudflare.com/workers-ai/models/whisper-large-v3-turbo/"),
        ("Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/"),
    ])

    base.add_heading(doc, "Vì sao Cloudflare rẻ?", 2)
    base.add_bullet(doc, "Workers AI có 10.000 Neurons/ngày miễn phí; Whisper Large V3 Turbo dùng 46,63 Neurons cho mỗi phút audio.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Đơn giá niêm yết của model là 0,00051 USD/phút audio; Cloudflare vẫn tính backend theo Neurons.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Không cần giữ GPU ở trạng thái sẵn sàng, nên chi phí tăng gần tuyến tính theo lượng audio thay vì theo giờ máy nhàn rỗi.", base.BULLET_NUM_ID)

    base.add_heading(doc, "Điểm cần kiểm soát", 2)
    base.add_body(doc, "Quota dùng chung: LLM, TTS hoặc dự án khác trong cùng tài khoản có thể ăn vào 10.000 Neurons/ngày.", bold_lead="Quota dùng chung:")
    base.add_body(doc, "Free plan dừng inference: Khi vượt quota, các thao tác tiếp theo lỗi cho đến khi quota reset; production từ 22 user nên bật Paid.", bold_lead="Free plan dừng inference:")
    base.add_body(doc, "Chất lượng: Phải đo WER/CER trên giọng nói thật, đặc biệt với tiếng Việt, giọng trẻ em, tiếng ồn và từ vựng chuyên ngành.", bold_lead="Chất lượng:")

    base.add_callout(
        doc,
        "NGÂN SÁCH KHUYẾN NGHỊ",
        "POC: Workers Free. Production 22-1.000 user: Workers Paid + ngân sách AI theo bảng. Đặt cảnh báo Neurons/ngày ở 70%, 85% và 95% để tránh bất ngờ.",
        fill=base.LIGHT_BLUE,
        accent=base.NAVY,
    )

    # Trang 3 - Vertex.
    base.add_page_break(doc)
    base.add_heading(doc, "2. Vertex AI custom endpoint", 1)
    base.add_body(
        doc,
        "Kịch bản này tự đóng gói Whisper hoặc model ASR tương đương vào custom container. Đây không phải Google Cloud Speech-to-Text, vốn là dịch vụ API riêng và có biểu phí khác.",
    )
    base.add_callout(
        doc,
        "ĐƠN GIÁ THAM CHIẾU",
        "n1-standard-4: 0,219 USD/giờ; NVIDIA Tesla T4: 0,402 USD/giờ. Tổng một replica: 0,621 USD/giờ. Một replica chạy 730 giờ/tháng = 453,33 USD/tháng.",
        fill=base.LIGHT_GRAY,
        accent=base.BLUE,
    )

    vertex_rows = []
    for users in SCALES:
        lo, hi = vertex_compute_range(users)
        if users == 1:
            scale_zero = f"{money(VERTEX_DAILY_FLOOR)}+"
        elif users == 100:
            scale_zero = f"{money(max(VERTEX_DAILY_FLOOR, lo))}-{money(hi)}"
        else:
            scale_zero = f"{money(lo)}-{money(hi)}"
        vertex_rows.append([
            f"{users:,}".replace(",", "."),
            f"{HOURS_PER_USER_MONTH * users:,.0f}".replace(",", "."),
            scale_zero,
            money(VERTEX_WARM_MONTH),
            money(VERTEX_HA_MONTH),
        ])
    base.add_table(
        doc,
        ["User", "Audio giờ/tháng", "Scale-to-zero*", "1 replica warm", "2 replica HA"],
        vertex_rows,
        [950, 1740, 2350, 2160, 2160],
        [WD_ALIGN_PARAGRAPH.CENTER] * 5,
        font_size=8.7,
        header_fill=base.PALE_GOLD,
    )
    add_caption_with_links(doc, [
        ("Google pricing", "https://cloud.google.com/products/gemini-enterprise-agent-platform/pricing"),
        ("Deploy-model / scale-to-zero", "https://docs.cloud.google.com/sdk/gcloud/reference/beta/ai/endpoints/deploy-model"),
    ])
    base.add_body(
        doc,
        "* Scale-to-zero là ước tính compute với tốc độ xử lý 10-30 lần realtime. Với tải hằng ngày, tài liệu ghi idle scaledown và min scale-up mặc định là 1 giờ khi min replica = 0; vì vậy 1 user được ngân sách hóa tối thiểu khoảng $18,63/tháng. Thực tế còn phụ thuộc cold start và phân bố phiên.",
        italic=True,
        color=base.MUTED,
    )

    base.add_heading(doc, "Chi phí không nằm trong bảng", 2)
    base.add_bullet(doc, "Cloud Storage cho model/audio, Artifact Registry, Cloud Logging, egress và thuế.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Thời gian MLOps: đóng gói container, CUDA/driver, autoscaling, benchmark, giám sát, rollback và vá bảo mật.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Số replica thực tế: một T4 chưa chắc đáp ứng 1.000 user nếu nhiều phiên đồng thời; phải load test trên traffic profile thật.", base.BULLET_NUM_ID)

    base.add_heading(doc, "Khi Vertex có thể thắng", 2)
    base.add_body(doc, "Nếu traffic được gom thành batch hoặc liên tục đủ lâu, GPU đạt utilization cao và model chạy nhanh hơn realtime, scale-to-zero có thể cạnh tranh về compute. Tuy nhiên với sản phẩm hội thoại, cold start và latency thường quan trọng hơn chi phí thuần.")

    # Trang 4 - difficulty, stop criteria, sources.
    base.add_page_break(doc)
    base.add_heading(doc, "3. Độ khó và vận hành", 1)
    base.add_table(
        doc,
        ["Tiêu chí", "Cloudflare Workers AI", "Vertex AI custom endpoint"],
        [
            ["Độ khó", "Thấp - khoảng 2/10", "Trung bình-cao - khoảng 7/10"],
            ["POC", "Khoảng 1-3 ngày", "Khoảng 2-4 tuần"],
            ["GPU/CUDA", "Không quản lý", "Phải cấu hình và benchmark"],
            ["Autoscaling", "Serverless", "Replica, cold start, scale-to-zero"],
            ["Model tùy biến", "Hạn chế theo catalog", "Cao; custom container/fine-tune"],
            ["IAM/VPC GCP", "Không native GCP", "Tích hợp sâu"],
            ["Rủi ro chính", "Quota và phụ thuộc catalog", "Chi phí nhàn rỗi, MLOps, latency"],
        ],
        [2450, 3200, 3710],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
        header_fill=base.LIGHT_BLUE,
    )

    base.add_heading(doc, "4. Khi nào dừng hoặc chuyển giải pháp?", 1)
    base.add_heading(doc, "Dừng Cloudflare-first khi", 2)
    base.add_bullet(doc, "WER/CER hoặc độ trễ không đạt SLA sau khi đã thử cấu hình hợp lý, và model custom cải thiện ít nhất 10-15% trên golden set.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Bắt buộc dữ liệu, IAM, VPC, audit log hoặc quy trình MLOps phải nằm trong Google Cloud.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Quy mô tiến gần 3.000 user, traffic đủ đều và benchmark cho thấy một T4 đáp ứng tải với TCO thấp hơn Cloudflare ít nhất 20%.", base.BULLET_NUM_ID)

    base.add_heading(doc, "Dừng Vertex warm khi", 2)
    base.add_bullet(doc, "Dưới khoảng 3.000 user mà không có lợi ích chất lượng/governance đủ lớn để bù chi phí và công vận hành.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "GPU utilization thấp, cold start làm hỏng trải nghiệm, hoặc cần tăng lên hai replica khiến điểm hòa vốn dịch lên khoảng 5.900 user.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Sau POC, tổng chi phí sở hữu không thấp hơn Cloudflare ít nhất 20% hoặc SLA không tốt hơn rõ rệt.", base.BULLET_NUM_ID)

    base.add_callout(
        doc,
        "LỘ TRÌNH ĐỀ XUẤT",
        "Bắt đầu Cloudflare; đo chi phí thực, WER/CER và p95 latency. Ở 500-1.000 user, chỉ chạy Vertex ở chế độ shadow nếu có mục tiêu chất lượng hoặc governance. Xem xét Vertex làm primary quanh 3.000+ user sau load test, không dựa riêng vào phép chia giá GPU.",
        fill=base.PALE_GREEN,
        accent=base.GREEN,
    )

    base.add_heading(doc, "Nguồn và giới hạn", 2)
    add_caption_with_links(doc, [
        ("Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/"),
        ("Whisper V3 Turbo", "https://developers.cloudflare.com/workers-ai/models/whisper-large-v3-turbo/"),
        ("Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/"),
        ("Google pricing", "https://cloud.google.com/products/gemini-enterprise-agent-platform/pricing"),
        ("Vertex deploy-model", "https://docs.cloud.google.com/sdk/gcloud/reference/beta/ai/endpoints/deploy-model"),
    ])
    base.add_body(doc, "Chưa gồm LLM, TTS, database, lưu audio, băng thông, thuế và nhân sự. Cần kiểm tra lại giá trước khi ký ngân sách.", italic=True, color=base.MUTED)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_report()

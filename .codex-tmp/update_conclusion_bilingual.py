from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SOURCE = Path(
    r"D:\Code\HuaMei\App_noi\be\9_10th8\speaking-ai-nextjs-backend\outputs\bao-cao-cloudflare-vs-vertex-ai-quy-mo-5000-user-viet-trung.docx"
)
OUTPUT = Path(
    r"D:\Code\HuaMei\App_noi\be\9_10th8\speaking-ai-nextjs-backend\outputs\bao-cao-cloudflare-vs-vertex-ai-viet-trung-cap-nhat-ket-luan.docx"
)

VI_1 = "Hiện tại, sử dụng Cloudflare làm giải pháp triển khai."
ZH_1 = "现阶段采用 Cloudflare 作为部署方案。"

VI_2 = (
    "Trong quá trình vận hành, tiếp tục thu thập dữ liệu giọng nói trẻ em, "
    "gồm file audio gốc, bản phiên âm/nhãn và thông tin mô tả; có thể sử dụng "
    "file Word để quản lý danh mục dữ liệu. Bộ dữ liệu này sẽ được bổ sung cho "
    "lớp nhận diện giọng trẻ em nhằm tăng độ chính xác, đồng thời có thể dùng "
    "để huấn luyện hoặc tinh chỉnh mô hình AI riêng trong tương lai."
)
ZH_2 = (
    "在运行过程中，继续采集儿童语音数据，包括原始音频、转写/标签及相关说明；"
    "可使用 Word 文件统一管理数据清单。这套数据将用于完善儿童语音识别模块，"
    "以提高识别准确率，并可在未来用于训练或微调自有 AI 模型。"
)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, size: float, color: str, chinese: bool = False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei" if chinese else "Arial")
    if chinese:
        rfonts.set(qn("w:hint"), "eastAsia")


def replace_bilingual(paragraph, vietnamese: str, chinese: str) -> None:
    clear_paragraph(paragraph)
    vi_run = paragraph.add_run(vietnamese)
    set_run_font(vi_run, 10.5, "1F2937")
    vi_run.add_break()
    zh_run = paragraph.add_run(chinese)
    set_run_font(zh_run, 8.5, "1F2937", chinese=True)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)
    if len(doc.paragraphs) < 9 or "KẾT LUẬN" not in doc.paragraphs[6].text:
        raise RuntimeError("Không tìm thấy phần kết luận theo cấu trúc mong đợi.")

    replace_bilingual(doc.paragraphs[7], VI_1, ZH_1)
    replace_bilingual(doc.paragraphs[8], VI_2, ZH_2)

    doc.core_properties.title = "Báo cáo Cloudflare và Vertex AI song ngữ Việt - Trung"
    doc.core_properties.subject = "Bản cập nhật kết luận và định hướng dữ liệu giọng nói trẻ em"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

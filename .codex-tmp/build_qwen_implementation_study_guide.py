from pathlib import Path
import importlib.util

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BASE_SCRIPT = ROOT / ".codex-tmp" / "build_qwen_training_roadmap_docx.py"
OUTPUT = ROOT / "deliverables" / "Noi_dung_can_tim_hieu_de_trien_khai_Qwen.docx"

spec = importlib.util.spec_from_file_location("qwen_doc_base", BASE_SCRIPT)
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
MUTED = "66758C"
GOLD = "B26A00"
INK = "20242A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"


def set_style_font(style, name="Calibri", size=11, color=INK, bold=False):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def configure_list_styles(doc):
    styles = doc.styles
    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        set_style_font(style, size=11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        base.set_run_font(r1, bold=True, color=DARK_BLUE)
        r2 = p.add_run(text[len(bold_prefix):])
        base.set_run_font(r2)
    else:
        r = p.add_run(text)
        base.set_run_font(r)
    return p


def create_numbering_instance(doc, style_name="List Number"):
    style = doc.styles[style_name]
    num_id_nodes = style._element.xpath("./w:pPr/w:numPr/w:numId")
    if not num_id_nodes:
        raise RuntimeError(f"Style {style_name} does not contain a numbering definition")
    base_num_id = int(num_id_nodes[0].get(qn("w:val")))
    numbering = doc.part.numbering_part.element
    base_num = numbering.xpath(f'./w:num[@w:numId="{base_num_id}"]')[0]
    abstract_id = base_num.xpath("./w:abstractNumId")[0].get(qn("w:val"))
    existing = [int(node.get(qn("w:numId"))) for node in numbering.xpath("./w:num")]
    new_num_id = max(existing) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_num_id


def add_number(doc, text, num_id):
    p = doc.add_paragraph(style="List Number")
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)
    r = p.add_run(text)
    base.set_run_font(r)
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        base.set_run_font(r1, bold=True, color=DARK_BLUE)
        r2 = p.add_run(text[len(bold_prefix):])
        base.set_run_font(r2)
    else:
        r = p.add_run(text)
        base.set_run_font(r)
    return p


def add_outcome(doc, lines):
    base.add_callout(
        doc,
        "Tiêu chí hoàn thành",
        " ".join(lines),
        fill=LIGHT_GRAY,
        accent=BLUE,
    )


def add_link_item(doc, label, url, description):
    p = doc.add_paragraph(style="List Bullet")
    base.add_hyperlink(p, label, url)
    r = p.add_run(f" - {description}")
    base.set_run_font(r)


def build_document():
    doc = Document()
    base.configure_styles(doc)
    configure_list_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("Cẩm nang học và triển khai Qwen | Ứng dụng Việt-Anh")
    base.set_run_font(hr, size=8.5, color=MUTED)

    fp = section.footer.paragraphs[0]
    base.add_page_number(fp)

    # Editorial cover.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    kr = kicker.add_run("CẨM NANG TRIỂN KHAI DÀNH CHO NGƯỜI MỚI")
    base.set_run_font(kr, size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    tr = title.add_run("SÁU NHÓM KIẾN THỨC CẦN HỌC\nĐỂ TRIỂN KHAI QWEN")
    base.set_run_font(tr, size=27, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run("Từ hiểu model, xây dữ liệu và QLoRA đến tích hợp Next.js production")
    base.set_run_font(sr, size=13.5, color=DARK_BLUE)

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.paragraph_format.space_after = Pt(46)
    dr = desc.add_run("Tài liệu học theo đầu ra: mỗi giai đoạn đều có kiến thức, bài thực hành và tiêu chí hoàn thành")
    base.set_run_font(dr, size=10.5, italic=True, color=MUTED)

    base.add_table(
        doc,
        ["Phạm vi", "Thời gian tham khảo", "Mục tiêu đầu tiên"],
        [["Qwen 4B cho dịch Việt-Anh", "6-10 tuần nếu đã biết lập trình", "Self-host và benchmark trước khi train"]],
        [3000, 3000, 3360],
        font_size=9.2,
    )

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.space_before = Pt(70)
    date_r = date_p.add_run("Phiên bản 1.0 | Tháng 8 năm 2026")
    base.set_run_font(date_r, size=10.5, bold=True, color=NAVY)

    doc.add_page_break()

    doc.add_heading("1. Cách sử dụng tài liệu", level=1)
    base.add_section_intro(doc, "Không cần học hết Machine Learning. Hãy học đủ để xây một vòng đời model có thể đo lường, triển khai và rollback an toàn.")
    base.add_callout(
        doc,
        "Nguyên tắc cốt lõi",
        "Chạy Qwen nguyên bản, tạo bộ test chuẩn và benchmark trước. Chỉ fine-tune khi đã xác định được nhóm lỗi lặp lại mà prompt hoặc rule không giải quyết ổn định.",
    )
    base.add_table(
        doc,
        ["Giai đoạn", "Sản phẩm phải tạo ra", "Câu hỏi cần trả lời"],
        [
            ["1. Model và hạ tầng", "Bảng capacity planning", "Cần GPU nào và chi phí bao nhiêu?"],
            ["2. Kiến thức nền", "Model chạy được qua REST API", "Đã đủ năng lực kỹ thuật chưa?"],
            ["3. Kho dữ liệu", "Golden set và dataset có version", "Dữ liệu có đúng và có quyền sử dụng không?"],
            ["4. Self-host", "Báo cáo benchmark", "Model gốc đã đủ tốt chưa?"],
            ["5. SFT + QLoRA", "Adapter và model card", "Fine-tune có tốt hơn baseline không?"],
            ["6. Next.js production", "Pilot có fallback", "Có vận hành và rollback an toàn không?"],
        ],
        [1850, 3450, 4060],
        font_size=8.9,
    )

    doc.add_heading("2. Tìm hiểu model, tham số và chi phí hạ tầng", level=1)
    doc.add_heading("2.1. Khái niệm model bắt buộc", level=2)
    for text in [
        "Base model và Instruct model: ứng dụng hội thoại nên bắt đầu từ bản Instruct.",
        "Token và tokenizer: model xử lý token số, không đọc trực tiếp từ hoặc câu.",
        "Chat template: quy tắc biến system, user và assistant thành chuỗi token.",
        "Context length: tổng token đầu vào và đầu ra mà request có thể sử dụng.",
        "Sampling: temperature, top_p, max_tokens và stop; dịch cần cấu hình ổn định, thường temperature rất thấp.",
        "Precision và quantization: FP32, BF16/FP16, INT8 và INT4.",
        "KV cache, throughput, latency và concurrency: bốn yếu tố quyết định dung lượng phục vụ thực tế.",
    ]:
        add_bullet(doc, text)

    base.add_callout(
        doc,
        "Model nên thử đầu tiên",
        "Qwen3-4B-Instruct-2507 là model non-thinking 4B, phù hợp để dựng baseline dịch câu ngắn. Ứng dụng không cần dùng context rất dài; bắt đầu ở khoảng 2.048-4.096 token để giảm KV cache và rủi ro OOM.",
    )

    doc.add_heading("2.2. Ước lượng VRAM", level=2)
    add_para(doc, "Bộ nhớ thô của trọng số được ước lượng bằng số tham số nhân số byte cho mỗi tham số. Runtime thực tế còn cần KV cache, CUDA, tensor trung gian và bộ nhớ dự phòng.")
    base.add_table(
        doc,
        ["Định dạng", "Trọng số thô của model 4B", "Ứng dụng"],
        [
            ["FP32", "Khoảng 16 GB", "Không phù hợp để inference tiết kiệm"],
            ["BF16 / FP16", "Khoảng 8 GB", "Chất lượng tốt, cần VRAM cao hơn"],
            ["INT8", "Khoảng 4 GB", "Giảm bộ nhớ"],
            ["INT4", "Khoảng 2 GB", "Phù hợp thử nghiệm và QLoRA"],
        ],
        [2200, 2700, 4460],
    )
    base.add_table(
        doc,
        ["Mục đích", "GPU/VRAM khởi đầu", "Lưu ý"],
        [
            ["Thử inference INT4", "6-8 GB", "Concurrency thấp"],
            ["Self-host context ngắn", "12-16 GB", "Benchmark P95 và KV cache"],
            ["Pilot QLoRA", "24 GB", "Điểm khởi đầu thực hành"],
            ["LLM + ASR + TTS", "24 GB trở lên", "Nên tách service và đo tải"],
        ],
        [2700, 2600, 4060],
    )

    doc.add_heading("2.3. Các nhóm chi phí", level=2)
    for text in [
        "GPU, CPU, RAM, SSD và băng thông.",
        "Điện, làm mát hoặc giá thuê GPU theo giờ.",
        "Lưu model, checkpoint, dataset, log và backup.",
        "Monitoring, bảo mật, thời gian kỹ sư và máy dự phòng.",
        "Chi phí do timeout, downtime, fallback và GPU nhàn rỗi.",
    ]:
        add_bullet(doc, text)
    add_outcome(doc, ["Có bảng GPU/VRAM, request đồng thời, độ dài prompt/output, token/giây, latency P50/P95, tỷ lệ OOM và chi phí trên 1.000 request."])

    doc.add_heading("3. Xác định kiến thức cần học và phần có thể bỏ qua", level=1)
    doc.add_heading("3.1. Nền tảng bắt buộc", level=2)
    base.add_table(
        doc,
        ["Nhóm", "Nội dung phải biết"],
        [
            ["Python", "Hàm, class, exception, JSON/JSONL/CSV, virtual environment và package"],
            ["Linux", "File, process, quyền, biến môi trường, log, CPU/RAM/GPU"],
            ["Git", "Branch, commit, tag và version code/dataset/config"],
            ["Docker", "Image, container, volume, port, health check và NVIDIA runtime"],
            ["HTTP/REST", "JSON, authentication, timeout, retry, streaming và status code"],
            ["AI", "Tokenizer, loss, learning rate, epoch, batch, split, overfitting, quantization"],
            ["Ứng dụng", "TypeScript, Next.js server-side, schema validation, log và feature flag"],
        ],
        [2100, 7260],
        font_size=9.1,
    )
    doc.add_heading("3.2. Có thể tạm bỏ qua", level=2)
    for text in [
        "Pretraining Transformer và tokenizer từ đầu.",
        "Tự viết CUDA kernel hoặc tối ưu compiler.",
        "Distributed training nhiều node, NVLink và InfiniBand.",
        "RLHF, reward model, DPO và GRPO.",
        "Mixture-of-Experts chuyên sâu.",
        "Tự train ASR, TTS hoặc voice cloning cùng lúc với LLM.",
    ]:
        add_bullet(doc, text)
    add_outcome(doc, ["Chạy được model Hugging Face bằng Python; dựng container GPU; gọi model qua REST; giải thích được token, context, loss và overfitting; đọc được log GPU và lỗi OOM."])

    doc.add_heading("4. Xây kho dữ liệu Việt-Anh có kiểm duyệt", level=1)
    base.add_section_intro(doc, "Chất lượng dữ liệu và test set quyết định giá trị của fine-tuning. Audio, transcript và bản dịch là ba loại dữ liệu khác nhau.")
    doc.add_heading("4.1. Dữ liệu cho từng model", level=2)
    base.add_table(
        doc,
        ["Model", "Đầu vào", "Nhãn chuẩn"],
        [
            ["ASR / Whisper", "Audio tiếng Việt", "Transcript tiếng Việt"],
            ["LLM / Qwen", "Câu tiếng Việt", "Bản dịch tiếng Anh"],
            ["TTS", "Văn bản tiếng Anh", "Audio tiếng Anh tương ứng"],
        ],
        [2400, 3300, 3660],
    )
    base.add_callout(doc, "Lưu ý", "Không đưa trực tiếp audio vào Qwen text model. Muốn train Qwen phải có transcript tiếng Việt và bản dịch tiếng Anh đã được con người xác nhận.")

    doc.add_heading("4.2. Schema và metadata", level=2)
    base.add_code(doc, '{\n  "id": "vi_en_000001",\n  "vietnamese": "Mẹ ơi, con không muốn mua cái này.",\n  "english": "Mom, I don\'t want to buy this.",\n  "context": "outside",\n  "region": "south",\n  "speech_act": "refusal",\n  "review_status": "human_approved",\n  "split": "train"\n}')
    for text in [
        "Lưu nguồn dữ liệu, reviewer, thời gian duyệt và phiên bản guideline.",
        "Dùng speaker_id đã ẩn danh thay cho tên thật.",
        "Gắn nhãn phủ định, chủ thể, người nghe, số lượng, màu sắc, thời gian và địa điểm.",
        "Lưu consent version và chính sách lưu/xóa dữ liệu nếu có dữ liệu trẻ em.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("4.3. Quy trình kiểm duyệt", level=2)
    review_num_id = create_numbering_instance(doc)
    for text in [
        "Người thứ nhất tạo hoặc sửa bản dịch.",
        "Người thứ hai kiểm tra độc lập.",
        "Bất đồng được chuyển cho người có quyền quyết định.",
        "Ghi lại lịch sử sửa và lý do sửa.",
        "Audit ngẫu nhiên 5-10% dữ liệu đã duyệt.",
    ]:
        add_number(doc, text, review_num_id)

    doc.add_heading("4.4. Chia tập và quy mô pilot", level=2)
    base.add_table(
        doc,
        ["Tập", "Tỷ lệ", "Quy tắc"],
        [
            ["Train", "80%", "Dùng cập nhật adapter"],
            ["Validation", "10%", "Chọn checkpoint và phát hiện overfitting"],
            ["Test", "10%", "Khóa kín, không dùng chọn tham số"],
        ],
        [2200, 1800, 5360],
    )
    for text in [
        "Chia audio theo người nói; không để cùng một trẻ xuất hiện ở train và test.",
        "Gom câu gần trùng theo cụm ngữ nghĩa trước khi chia.",
        "Bắt đầu với 500-1.000 mẫu golden test và 2.000-5.000 mẫu SFT có human review.",
        "Không dùng output Cloudflare chưa kiểm duyệt làm đáp án chuẩn.",
    ]:
        add_bullet(doc, text)
    add_outcome(doc, ["Có guideline, dataset được version, báo cáo chất lượng, golden test đã khóa và danh sách nguồn/consent."])

    doc.add_heading("5. Chạy Qwen self-host và benchmark trước khi train", level=1)
    doc.add_heading("5.1. Kiến trúc pilot", level=2)
    base.add_code(doc, "Next.js backend\n    ↓ HTTPS + API key\nQwen model service riêng\n    ↓\nvLLM + GPU + Qwen3-4B-Instruct-2507")
    for text in [
        "Không nạp model trực tiếp vào tiến trình Next.js hoặc serverless function.",
        "Khóa revision model, tokenizer, vLLM, PyTorch, CUDA và Docker image.",
        "Bắt đầu max-model-len khoảng 4.096 cho câu dịch ngắn.",
        "Đặt temperature rất thấp và giới hạn max_tokens.",
        "Bảo vệ endpoint bằng HTTPS, API key, firewall và network policy.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("5.2. Ma trận benchmark", level=2)
    base.add_table(
        doc,
        ["Ứng viên", "Mục đích", "Điều kiện đánh giá"],
        [
            ["Cloudflare", "Baseline hiện tại", "Cùng prompt và cùng golden test"],
            ["Qwen nguyên bản", "Xác định có cần train", "Không kém đáng kể ở lỗi trọng yếu"],
            ["Qwen + prompt/rule", "Tối ưu không train", "Ổn định hơn model gốc"],
            ["Qwen + LoRA", "Ứng viên production", "Tốt hơn baseline mà không overfit"],
        ],
        [2300, 3100, 3960],
        font_size=9.0,
    )
    doc.add_heading("5.3. Chỉ số phải đo", level=2)
    for text in [
        "Giữ đúng phủ định, người nói/người nghe và quan hệ gia đình.",
        "Không mất màu sắc, số lượng, thời gian hoặc địa điểm.",
        "Không tự thêm nghĩa, lịch sự hoặc hoàn chỉnh câu thiếu.",
        "Đúng định dạng: chỉ trả một câu tiếng Anh nếu ứng dụng yêu cầu.",
        "Human acceptance rate, latency P50/P95, token/giây, timeout, OOM và chi phí.",
        "Load test ở nhiều mức concurrency và độ dài request.",
    ]:
        add_bullet(doc, text)
    base.add_callout(doc, "Quyết định go/no-go", "Chỉ fine-tune khi lỗi lặp lại, prompt/rule không giải quyết ổn định, có đủ ví dụ chất lượng và golden test đo được mức cải thiện.")
    add_outcome(doc, ["Có model server OpenAI-compatible, báo cáo Cloudflare-vs-Qwen, danh sách nhóm lỗi và quyết định có cần fine-tune hay không."])

    doc.add_heading("6. Fine-tune bằng SFT + QLoRA", level=1)
    doc.add_heading("6.1. Cơ chế", level=2)
    base.add_code(doc, "Qwen base nạp 4-bit và được đóng băng\n                +\nLoRA adapter nhỏ được train ở precision cao hơn\n                ↓\nBase model + adapter = model theo miền ứng dụng")
    for text in [
        "SFT học từ cặp câu đầu vào và câu trả lời chuẩn.",
        "LoRA chỉ cập nhật ma trận low-rank, giữ nguyên trọng số gốc.",
        "QLoRA nạp base model ở 4-bit để giảm VRAM.",
        "Adapter nhỏ, dễ lưu version, phân phối và rollback.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("6.2. Định dạng hội thoại", level=2)
    base.add_code(doc, '{\n  "messages": [\n    {"role": "system", "content": "Translate Vietnamese to English faithfully."},\n    {"role": "user", "content": "Con không muốn đi học."},\n    {"role": "assistant", "content": "I don\'t want to go to school."}\n  ]\n}')
    base.add_callout(doc, "Loss", "Ưu tiên tính loss trên phần assistant/completion để model học câu trả lời chuẩn, không học lại toàn bộ prompt. Phải kiểm tra chat template tương thích.")

    doc.add_heading("6.3. Tham số pilot tham khảo", level=2)
    base.add_table(
        doc,
        ["Tham số", "Điểm khởi đầu", "Mục đích"],
        [
            ["Quantization", "NF4 4-bit", "Giảm VRAM"],
            ["LoRA rank", "16 hoặc 32", "Bắt đầu nhỏ rồi benchmark"],
            ["LoRA alpha", "16-32", "Điều khiển mức scaling"],
            ["LoRA dropout", "Khoảng 0,05", "Hạn chế overfitting"],
            ["Sequence length", "512-1.024", "Phù hợp câu ngắn"],
            ["Epoch", "2-3", "Dừng sớm nếu validation xấu đi"],
            ["Effective batch", "16-32", "Dùng gradient accumulation"],
            ["Compute dtype", "BF16 nếu hỗ trợ", "Ổn định và hiệu quả"],
        ],
        [2400, 2500, 4460],
        font_size=8.9,
    )

    doc.add_heading("6.4. Trình tự một lần train", level=2)
    train_num_id = create_numbering_instance(doc)
    for text in [
        "Khóa version base model, tokenizer, dataset và training config.",
        "Kiểm tra schema, trùng lặp, PII và split.",
        "Chạy baseline trên test set trước khi train.",
        "Train adapter và theo dõi training/validation loss.",
        "Chọn checkpoint theo validation, không chọn theo cảm giác.",
        "Chạy test kín đúng một quy trình đã định trước.",
        "Human review và phân tích theo từng nhóm lỗi.",
        "Lưu adapter, model card, config và báo cáo.",
    ]:
        add_number(doc, text, train_num_id)
    add_outcome(doc, ["Có adapter, biểu đồ loss, dataset manifest, training config và báo cáo chứng minh model tốt hơn baseline mà không overfit."])

    doc.add_heading("7. Tích hợp backend Next.js và triển khai production", level=1)
    doc.add_heading("7.1. Provider abstraction", level=2)
    base.add_code(doc, "translateVietnameseToEnglish(input, options)\n    ├─ cloudflare\n    └─ self_hosted → OpenAI-compatible API")
    for text in [
        "Giữ interface chung để phần còn lại của ứng dụng không phụ thuộc trực tiếp vLLM.",
        "Tách provider, prompt builder, output normalization và telemetry.",
        "Thêm feature flag để đổi provider mà không thay toàn bộ luồng ứng dụng.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("7.2. Biến môi trường và bảo mật", level=2)
    base.add_code(doc, "TEXT_AI_PROVIDER=self_hosted\nSELF_HOSTED_LLM_BASE_URL=https://gpu.example.com/v1\nSELF_HOSTED_LLM_API_KEY=<secret>\nSELF_HOSTED_LLM_MODEL=Qwen/Qwen3-4B-Instruct-2507\nSELF_HOSTED_LLM_TIMEOUT_MS=2500")
    for text in [
        "API key chỉ tồn tại server-side; không dùng tiền tố NEXT_PUBLIC_.",
        "Endpoint GPU dùng HTTPS, firewall, allowlist và quyền truy cập tối thiểu.",
        "Không log nguyên văn dữ liệu nhạy cảm; log request ID và metadata vận hành.",
        "Xác thực người dùng trước Route Handler nếu endpoint có thể bị gọi từ bên ngoài.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("7.3. Timeout, fallback và kiểm soát đầu ra", level=2)
    base.add_code(doc, "Cache → Qwen self-host → kiểm tra schema/fidelity\n              ↓ timeout, OOM hoặc output lỗi\n              Cloudflare fallback")
    for text in [
        "Giới hạn input, max_tokens, timeout và request size.",
        "Chỉ retry tối đa một lần có kiểm soát.",
        "Dùng queue, concurrency limit và circuit breaker.",
        "Xác thực schema, độ dài và nội dung bắt buộc trước khi trả cho người dùng.",
        "Ghi model version, adapter version, latency và fallback reason.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("7.4. Rollout", level=2)
    base.add_code(doc, "Offline benchmark → Shadow traffic → Canary 5% → 25% → 50% → 100%")
    for text in [
        "Shadow: Qwen chạy song song nhưng Cloudflare vẫn trả kết quả.",
        "Canary: chỉ một tỷ lệ nhỏ request dùng Qwen.",
        "Tăng tỷ lệ khi chất lượng, latency và error rate đạt ngưỡng.",
        "Giữ feature flag và rollback adapter/model trong suốt rollout.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("7.5. Monitoring và kiểm thử", level=2)
    base.add_table(
        doc,
        ["Nhóm", "Cần theo dõi/kiểm thử"],
        [
            ["Chất lượng", "Acceptance rate, lỗi phủ định, fidelity và output không hợp lệ"],
            ["Hiệu năng", "P50/P95/P99, token/giây, queue depth và concurrency"],
            ["GPU", "Utilization, VRAM, KV cache, OOM và nhiệt độ"],
            ["Độ tin cậy", "Timeout, fallback, health/readiness và circuit breaker"],
            ["Kiểm thử", "Unit, contract, integration, load, failure injection và end-to-end"],
        ],
        [2300, 7060],
        font_size=9.1,
    )
    add_outcome(doc, ["Có pilot production với provider self-host, feature flag, fallback, dashboard, cảnh báo và quy trình rollback đã được diễn tập."])

    doc.add_heading("8. Lộ trình thực hiện đề xuất", level=1)
    base.add_table(
        doc,
        ["Tuần", "Trọng tâm", "Đầu ra"],
        [
            ["1", "Model, token, GPU và chi phí", "Bảng capacity planning"],
            ["2", "Python, Linux, Docker và REST", "Model chạy qua API"],
            ["3-4", "Golden dataset và guideline", "500-1.000 mẫu đã duyệt"],
            ["5", "Self-host và benchmark", "Báo cáo go/no-go cho fine-tune"],
            ["6-7", "SFT + QLoRA", "Adapter và model card"],
            ["8-10", "Next.js, shadow, canary", "Pilot production có fallback"],
        ],
        [1200, 3600, 4560],
    )
    base.add_callout(doc, "Bước tiếp theo", "Dựng Qwen nguyên bản và benchmark trên 500 câu chuẩn. Đây là bước tạo nhiều thông tin nhất với chi phí và rủi ro thấp hơn so với train ngay.")

    doc.add_heading("9. Tài liệu học chính thức", level=1)
    resources = [
        ("Qwen3-4B-Instruct-2507", "https://huggingface.co/Qwen/Qwen3-4B-Instruct-2507", "model card và hướng dẫn chạy"),
        ("Qwen với vLLM", "https://qwen.readthedocs.io/en/stable/deployment/vllm.html", "OpenAI-compatible model server"),
        ("Hugging Face LLM Course", "https://huggingface.co/learn/llm-course/en/chapter1/1", "tokenizer, Transformer và fine-tuning"),
        ("TRL SFTTrainer", "https://huggingface.co/docs/trl/sft_trainer", "SFT dataset hội thoại"),
        ("TRL PEFT/QLoRA", "https://huggingface.co/docs/trl/en/peft_integration", "4-bit quantization và adapter"),
        ("PEFT LoRA", "https://huggingface.co/docs/peft/main/conceptual_guides/lora", "cơ chế và tham số LoRA"),
        ("Next.js production checklist", "https://nextjs.org/docs/app/guides/production-checklist", "bảo mật và vận hành production"),
    ]
    for label, url, description in resources:
        add_link_item(doc, label, url, description)

    note = doc.add_paragraph(style="Caption")
    nr = note.add_run("Ghi chú: Các con số VRAM, thời gian và hyperparameter là điểm khởi đầu thực hành. Kết quả cuối cùng phải được benchmark trên GPU, dữ liệu và mức concurrency thực tế của dự án.")
    base.set_run_font(nr, size=9, italic=True, color=MUTED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

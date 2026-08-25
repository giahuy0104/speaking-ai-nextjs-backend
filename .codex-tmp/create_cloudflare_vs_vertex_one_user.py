from __future__ import annotations

import importlib.util
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
BASE_PATH = ROOT / ".codex-tmp" / "create_cloudflare_vs_selfhost_report.py"
OUTPUT_DIR = ROOT / "outputs"
OUTPUT_PATH = OUTPUT_DIR / "bao-cao-chi-phi-cloudflare-vs-vertex-ai-1-user.docx"

spec = importlib.util.spec_from_file_location("base_report", BASE_PATH)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)


USD_VND = 26_000
AUDIO_MINUTES_PER_DAY = 10
DAYS_PER_MONTH = 30
AUDIO_MINUTES_PER_MONTH = AUDIO_MINUTES_PER_DAY * DAYS_PER_MONTH

CLOUDFLARE_PRICE_PER_MINUTE = 0.00051
CLOUDFLARE_NEURONS_PER_MINUTE = 46.63
CLOUDFLARE_FREE_NEURONS_PER_DAY = 10_000
CLOUDFLARE_LIST_VALUE = AUDIO_MINUTES_PER_MONTH * CLOUDFLARE_PRICE_PER_MINUTE
CLOUDFLARE_NEURONS_PER_DAY = AUDIO_MINUTES_PER_DAY * CLOUDFLARE_NEURONS_PER_MINUTE

VERTEX_MACHINE_HOUR = 0.219
VERTEX_T4_HOUR = 0.402
VERTEX_NODE_HOUR = VERTEX_MACHINE_HOUR + VERTEX_T4_HOUR
VERTEX_WARM_MONTH = VERTEX_NODE_HOUR * 730
VERTEX_DEFAULT_SCALE_ZERO_MONTH = VERTEX_NODE_HOUR * 30
VERTEX_PURE_COMPUTE_LOW = (AUDIO_MINUTES_PER_MONTH / 60 / 30) * VERTEX_NODE_HOUR
VERTEX_PURE_COMPUTE_HIGH = (AUDIO_MINUTES_PER_MONTH / 60 / 10) * VERTEX_NODE_HOUR


def add_caption_with_links(doc, items: list[tuple[str, str]]) -> None:
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.space_before = base.Pt(4)
    p.paragraph_format.space_after = base.Pt(4)
    p.add_run("Nguồn: ")
    for idx, (label, url) in enumerate(items):
        if idx:
            p.add_run("; ")
        base.add_hyperlink(p, label, url)
    p.add_run(".")


def add_metadata(doc, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = base.Pt(3)
    base.set_run_font(p.add_run(f"{label}: "), size=10.5, color=base.NAVY, bold=True)
    base.set_run_font(p.add_run(value), size=10.5, color=base.INK)


def build_report() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    base.configure_styles(doc)
    base.configure_sections(doc)
    base.BULLET_NUM_ID, base.NUMBER_NUM_ID = base.add_custom_numbering(doc)

    props = doc.core_properties
    props.title = "Chi phí Cloudflare và Vertex AI cho một người dùng"
    props.subject = "Một người dùng sử dụng ASR 10 phút mỗi ngày"
    props.author = "Speaking AI"
    props.last_modified_by = "Speaking AI"
    props.keywords = "Cloudflare Workers AI, Vertex AI, Whisper, T4, ASR, chi phí"
    props.created = datetime(2026, 8, 21, tzinfo=timezone.utc)

    # Trang 1: Kết quả chính.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = base.Pt(18)
    p.paragraph_format.space_after = base.Pt(4)
    base.set_run_font(p.add_run("BÁO CÁO CHI PHÍ ASR"), size=10.5, color=base.BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = base.Pt(6)
    base.set_run_font(p.add_run("Cloudflare hay Vertex AI\ncho 1 người dùng?"), size=25, color=base.NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = base.Pt(16)
    base.set_run_font(p.add_run("Giả định sử dụng nhận dạng giọng nói 10 phút mỗi ngày"), size=14, color=base.MUTED)

    add_metadata(doc, "Khối lượng", "10 phút/ngày x 30 ngày = 300 phút/tháng = 5 giờ audio")
    add_metadata(doc, "Vertex tham chiếu", "Custom Whisper endpoint, n1-standard-4 + NVIDIA Tesla T4, us-central1")
    add_metadata(doc, "Ngày giá", "21/08/2026; quy đổi 26.000 VND/USD; chưa gồm thuế")

    base.add_callout(
        doc,
        "KẾT LUẬN",
        "Chọn Cloudflare. Với 10 phút/ngày, Workers AI chỉ dùng khoảng 466 Neurons/ngày, thấp hơn quota miễn phí 10.000 Neurons/ngày. Chi phí Cloudflare có thể là 0 USD/tháng trên Free plan hoặc 5 USD/tháng nếu chủ động dùng Workers Paid. Vertex AI không kinh tế cho một người dùng.",
        fill=base.PALE_GREEN,
        accent=base.GREEN,
    )

    base.add_heading(doc, "Kết quả tính nhanh", 1)
    base.add_table(
        doc,
        ["Phương án", "Chi phí tháng", "Quy đổi VND", "Đánh giá"],
        [
            ["Cloudflare Free", "$0", "0đ", "Tốt nhất; nằm trong quota AI miễn phí"],
            ["Cloudflare Paid", "$5,00", "130.000đ", "Production nhỏ; ASR vẫn trong quota miễn phí"],
            ["Vertex scale-to-zero", "Khoảng $18,63+", "Khoảng 484.000đ+", "Có cold start; ước tính một giờ hoạt động/ngày"],
            ["Vertex luôn chạy", "$453,33", "11,79 triệu", "Độ trễ tốt hơn nhưng quá đắt cho 1 user"],
        ],
        [2300, 1700, 1900, 3460],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.9,
        header_fill=base.LIGHT_BLUE,
    )
    base.add_body(
        doc,
        "Giá trị suy theo đơn giá Cloudflare là 300 x 0,00051 = 0,153 USD/tháng (khoảng 3.978đ), nhưng khoản này không bị tính tiền khi mỗi ngày vẫn dưới quota Neurons miễn phí.",
        italic=True,
        color=base.MUTED,
    )

    base.add_heading(doc, "Quyết định thực tế", 1)
    base.add_bullet(doc, "POC hoặc một người dùng: Cloudflare Free, tổng hạ tầng ASR có thể bằng 0.", base.BULLET_NUM_ID, "POC hoặc một người dùng:")
    base.add_bullet(doc, "Production nhỏ: Cloudflare Paid 5 USD/tháng nếu cần hạn mức Workers cao hơn và tránh request bị chặn khi vượt Free plan.", base.BULLET_NUM_ID, "Production nhỏ:")
    base.add_bullet(doc, "Vertex AI: chỉ dùng khi cần model tùy biến, IAM/VPC của GCP hoặc yêu cầu dữ liệu đặc thù; không chọn để tiết kiệm ở quy mô này.", base.BULLET_NUM_ID, "Vertex AI:")

    # Trang 2: Cloudflare.
    base.add_page_break(doc)
    base.add_heading(doc, "1. Tính chi phí Cloudflare", 1)
    base.add_callout(
        doc,
        "CÔNG THỨC SỬ DỤNG",
        "10 phút/ngày x 30 ngày = 300 phút/tháng. Mỗi ngày: 10 x 46,63 = 466,3 Neurons. Tỷ lệ dùng quota: 466,3 / 10.000 = 4,66%.",
        fill=base.LIGHT_GRAY,
        accent=base.BLUE,
    )
    base.add_table(
        doc,
        ["Chỉ số", "Kết quả", "Ý nghĩa"],
        [
            ["Audio mỗi ngày", "10 phút", "Khối lượng người dùng yêu cầu"],
            ["Audio mỗi tháng", "300 phút", "Tương đương 5 giờ audio"],
            ["Neurons mỗi ngày", "466,3", "Chỉ dùng 4,66% quota miễn phí"],
            ["Quota miễn phí", "10.000/ngày", "Reset hằng ngày lúc 00:00 UTC"],
            ["Giá trị theo list price", "$0,153/tháng", "Không bị tính nếu không vượt quota ngày"],
            ["Chi phí Free plan", "$0/tháng", "Giả định không có tải Workers AI khác"],
            ["Chi phí Paid plan", "$5/tháng", "Mức tối thiểu của Workers Paid"],
        ],
        [2600, 2200, 4560],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9.2,
        header_fill=base.PALE_GREEN,
    )
    add_caption_with_links(doc, [
        ("Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/"),
        ("Whisper Large V3 Turbo", "https://developers.cloudflare.com/workers-ai/models/whisper-large-v3-turbo/"),
        ("Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/"),
    ])

    base.add_heading(doc, "Cloudflare Free có đủ không?", 2)
    base.add_body(
        doc,
        "Có, nếu tài khoản chỉ phục vụ người dùng này và tổng tải Workers AI mỗi ngày không vượt 10.000 Neurons. Với 46,63 Neurons/phút, quota tương đương khoảng 214 phút Whisper Large V3 Turbo mỗi ngày; nhu cầu 10 phút/ngày chỉ chiếm khoảng 1/21 quota.",
    )
    base.add_bullet(doc, "Workers Free cho phép 100.000 request/ngày; một người dùng thông thường thấp hơn rất nhiều.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Nếu dùng chung tài khoản cho LLM, TTS hoặc người dùng khác, quota 10.000 Neurons/ngày được chia sẻ và có thể hết sớm hơn.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Khi vượt quota trên Free plan, inference tiếp theo sẽ lỗi; muốn tiếp tục cần nâng Workers Paid và trả phần vượt mức.", base.BULLET_NUM_ID)

    base.add_callout(
        doc,
        "NGÂN SÁCH ĐỀ XUẤT",
        "Giai đoạn thử nghiệm: 0 USD/tháng. Giai đoạn production nhỏ: dự trù 5 USD/tháng để có Workers Paid, dù riêng phần ASR của một user vẫn nằm trong quota AI miễn phí.",
        fill=base.LIGHT_BLUE,
        accent=base.NAVY,
    )

    # Trang 3: Vertex và so sánh.
    base.add_page_break(doc)
    base.add_heading(doc, "2. Tính chi phí Vertex AI custom endpoint", 1)
    base.add_body(
        doc,
        "Đây là kịch bản tự đóng gói Whisper hoặc model tương đương vào custom container và triển khai lên Vertex AI Endpoint. Nó không phải Google Speech-to-Text, vốn là một API riêng tính phí theo phút.",
    )
    base.add_callout(
        doc,
        "ĐƠN GIÁ THAM CHIẾU",
        "n1-standard-4: 0,219 USD/giờ; NVIDIA Tesla T4: 0,402 USD/giờ. Tổng một replica: 0,621 USD/giờ tại us-central1.",
        fill=base.LIGHT_GRAY,
        accent=base.BLUE,
    )

    base.add_heading(doc, "2.1. Endpoint luôn sẵn sàng", 2)
    base.add_body(doc, "0,621 x 730 giờ/tháng = 453,33 USD/tháng, tương đương khoảng 11,79 triệu đồng. Chỉ có một user nên gần như toàn bộ GPU sẽ nhàn rỗi.")

    base.add_heading(doc, "2.2. Scale-to-zero", 2)
    base.add_body(
        doc,
        "Compute thuần cho 5 giờ audio/tháng có thể chỉ khoảng 0,10-0,31 USD nếu GPU xử lý nhanh hơn realtime 10-30 lần. Tuy nhiên con số này không phản ánh cách endpoint tương tác thực tế.",
    )
    base.add_bullet(doc, "Vertex cho phép min replica bằng 0, nhưng scale-to-zero có cold start khi model phải tải lại.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Tài liệu deploy-model ghi min scale-up mặc định là một giờ khi min replica bằng 0.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Nếu người dùng phát sinh một phiên riêng mỗi ngày, ước tính thực tế tối thiểu khoảng 30 giờ x 0,621 = 18,63 USD/tháng, chưa gồm startup, log, storage và egress.", base.BULLET_NUM_ID)

    base.add_table(
        doc,
        ["Kịch bản Vertex", "Chi phí compute", "Độ trễ", "Phù hợp 1 user?"],
        [
            ["Compute thuần lý thuyết", "$0,10-$0,31", "Không phản ánh cold start", "Không dùng làm ngân sách production"],
            ["Scale-to-zero, dùng mỗi ngày", "Khoảng $18,63+", "Cold start cao", "Không khuyến nghị cho hội thoại"],
            ["Một replica luôn chạy", "$453,33", "Thấp và ổn định", "Không kinh tế"],
            ["Hai replica HA", "$906,66", "Tốt hơn về HA", "Hoàn toàn không phù hợp"],
        ],
        [2400, 1900, 2100, 2960],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
        header_fill=base.PALE_GOLD,
    )
    add_caption_with_links(doc, [
        ("Vertex/Agent Platform pricing", "https://cloud.google.com/products/gemini-enterprise-agent-platform/pricing"),
        ("Vertex deploy-model scale-to-zero", "https://docs.cloud.google.com/sdk/gcloud/reference/beta/ai/endpoints/deploy-model"),
    ])
    base.add_heading(doc, "Nguồn kỹ thuật", 2)
    base.add_source(doc, "Cloudflare Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/", "10.000 Neurons/ngày miễn phí và đơn giá model")
    base.add_source(doc, "Cloudflare Whisper Large V3 Turbo", "https://developers.cloudflare.com/workers-ai/models/whisper-large-v3-turbo/", "0,00051 USD/phút audio")
    base.add_source(doc, "Cloudflare Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/", "Free plan và Paid plan tối thiểu 5 USD")
    base.add_source(doc, "Vertex/Agent Platform pricing", "https://cloud.google.com/products/gemini-enterprise-agent-platform/pricing", "giá n1-standard-4 và T4 theo giờ")
    base.add_source(doc, "Vertex deploy-model", "https://docs.cloud.google.com/sdk/gcloud/reference/beta/ai/endpoints/deploy-model", "min replica 0 và scale-to-zero")

    # Trang 4: So sánh và khuyến nghị.
    base.add_page_break(doc)
    base.add_heading(doc, "3. So sánh Cloudflare và Vertex AI", 1)
    base.add_table(
        doc,
        ["Tiêu chí", "Cloudflare Workers AI", "Vertex AI custom endpoint"],
        [
            ["Chi phí 1 user", "$0 Free hoặc $5 Paid", "$18,63+ scale-to-zero; $453,33 warm"],
            ["Độ khó", "Thấp - khoảng 2/10", "Trung bình-cao - khoảng 7/10"],
            ["Thời gian POC", "Khoảng 1-3 ngày", "Khoảng 2-4 tuần"],
            ["Quản lý GPU/CUDA", "Không", "Phải chọn máy, GPU và container"],
            ["Cold start", "Thấp", "Có khi scale-to-zero"],
            ["Model tùy biến", "Hạn chế", "Cao; có thể fine-tune model riêng"],
            ["IAM/VPC GCP", "Không nằm trong GCP", "Tích hợp tốt với hệ sinh thái GCP"],
            ["Vận hành", "Cloudflare quản lý", "Vẫn cần MLOps, benchmark và monitoring"],
        ],
        [2380, 3230, 3750],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
        header_fill=base.LIGHT_BLUE,
    )

    base.add_heading(doc, "Khi nào nên chọn Vertex dù chi phí cao hơn?", 2)
    base.add_body(doc, "Chất lượng: Model Cloudflare không đạt yêu cầu cho giọng trẻ em hoặc từ vựng chuyên biệt, và model tùy biến chứng minh cải thiện ít nhất 10-15% trên golden set.", bold_lead="Chất lượng:")
    base.add_body(doc, "Governance: Tổ chức bắt buộc dùng IAM, VPC, audit log, data residency hoặc MLOps trong Google Cloud.", bold_lead="Governance:")
    base.add_body(doc, "Quy mô: Hệ thống tăng lên hàng nghìn user, GPU có utilization tốt và TCO của Vertex thấp hơn Cloudflare ít nhất 20% sau benchmark thực tế.", bold_lead="Quy mô:")

    base.add_heading(doc, "Quyết định đề xuất", 2)
    base.add_callout(
        doc,
        "CHỌN CLOUDFLARE",
        "Với một người dùng nói 10 phút/ngày, Cloudflare rẻ hơn từ khoảng 3,7 lần so với Vertex scale-to-zero thực tế và khoảng 91 lần so với một Vertex replica luôn chạy, nếu so với gói Cloudflare Paid 5 USD. Nếu dùng Cloudflare Free, chênh lệch còn lớn hơn.",
        fill=base.PALE_GREEN,
        accent=base.GREEN,
    )
    base.add_numbered(doc, "Bắt đầu bằng Workers Free và theo dõi Neurons/ngày.", base.NUMBER_NUM_ID)
    base.add_numbered(doc, "Nâng Workers Paid khi cần production ổn định hoặc có thêm người dùng/model dùng chung quota.", base.NUMBER_NUM_ID)
    base.add_numbered(doc, "Chỉ tạo Vertex POC ở chế độ shadow khi có mục tiêu chất lượng hoặc governance cụ thể.", base.NUMBER_NUM_ID)

    base.add_heading(doc, "4. Giới hạn của ước tính", 1)
    base.add_body(
        doc,
        "Báo cáo chỉ tính ASR. Chưa gồm model hội thoại, dịch văn bản, TTS, database, lưu audio, băng thông, thuế và nhân sự. Chi phí Vertex thay đổi theo region, machine type, accelerator, thời gian cold start và cách phân bố request trong ngày.",
    )

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_report()

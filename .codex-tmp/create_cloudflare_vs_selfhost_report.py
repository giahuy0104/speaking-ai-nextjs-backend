from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs"
TMP_DIR = ROOT / ".codex-tmp" / "cloudflare-selfhost-report"
OUTPUT_PATH = OUTPUT_DIR / "bao-cao-so-sanh-cloudflare-va-ai-tu-van-hanh.docx"
CHART_PATH = TMP_DIR / "chi-phi-asr-cloudflare-vs-gpu.png"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EAF5EA"
GREEN = "276749"
PALE_GOLD = "FFF7E0"
GOLD = "7A5A00"
PALE_RED = "FDECEC"
RED = "9B1C1C"
WHITE = "FFFFFF"
BORDER = "C9D2DC"

USD_VND = 26_000


def set_run_font(run, name="Arial", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = trpr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        trpr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    cant_split = trpr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        trpr.append(cant_split)


def set_table_borders(table, color=BORDER, size=6):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(indent_dxa))
    tblind.set(qn("w:type"), "dxa")

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    set_table_borders(table)


def style_table(table, header_fill=LIGHT_GRAY, font_size=9.3, alignments=None, zebra=True):
    for ridx, row in enumerate(table.rows):
        if ridx == 0:
            set_repeat_table_header(row)
        for cidx, cell in enumerate(row.cells):
            if ridx == 0:
                set_cell_shading(cell, header_fill)
            elif zebra and ridx % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if alignments and cidx < len(alignments):
                    paragraph.alignment = alignments[cidx]
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        color=NAVY if ridx == 0 else INK,
                        bold=True if ridx == 0 else None,
                    )


def add_table(doc, headers, rows, widths_dxa, alignments=None, font_size=9.3, header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, text in enumerate(headers):
        table.rows[0].cells[idx].text = text
    for row_data in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            cells[idx].text = str(text)
    set_table_geometry(table, widths_dxa)
    style_table(table, header_fill=header_fill, font_size=font_size, alignments=alignments)
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.extend([rfonts, color, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([rpr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    rpr.extend([color, size])
    r.extend([rpr, begin, instr, separate, text, end])
    paragraph._p.append(r)


def add_custom_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abstract = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]

    def make_numbering(num_fmt, lvl_text, font="Arial"):
        abstract_id = max(existing_abstract + [0]) + 1
        existing_abstract.append(abstract_id)
        num_id = max(existing_num + [0]) + 1
        existing_num.append(num_id)

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        txt = OxmlElement("w:lvlText")
        txt.set(qn("w:val"), lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.extend([tabs, ind, spacing])
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), font)
        rfonts.set(qn("w:hAnsi"), font)
        rpr.append(rfonts)
        lvl.extend([start, fmt, txt, jc, ppr, rpr])
        abstract.append(lvl)
        first_num_index = next(
            (idx for idx, child in enumerate(numbering) if child.tag == qn("w:num")),
            len(numbering),
        )
        numbering.insert(first_num_index, abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(abstract_id))
        num.append(abs_id)
        numbering.append(num)
        return num_id

    return make_numbering("bullet", "•"), make_numbering("decimal", "%1.")


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, numid])


def add_bullet(doc, text, bullet_num_id, bold_lead=None):
    p = doc.add_paragraph()
    apply_num(p, bullet_num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_numbered(doc, text, number_num_id, bold_lead=None):
    p = doc.add_paragraph()
    apply_num(p, number_num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_body(doc, text="", bold_lead=None, italic=False, color=INK, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, italic=italic, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic, color=color)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.083)
    p.paragraph_format.right_indent = Inches(0.083)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "6")
        border.set(qn("w:color"), accent)
        borders.append(border)
    ppr.append(borders)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def create_chart(path):
    width, height = 1400, 720
    img = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(img)
    font_regular = "C:/Windows/Fonts/arial.ttf"
    font_bold = "C:/Windows/Fonts/arialbd.ttf"
    title_font = ImageFont.truetype(font_bold, 42)
    label_font = ImageFont.truetype(font_regular, 27)
    small_font = ImageFont.truetype(font_regular, 23)
    bold_font = ImageFont.truetype(font_bold, 25)

    draw.text((80, 35), "Chi phí ASR mỗi tháng: Cloudflare và một GPU cloud luôn chạy", font=title_font, fill=f"#{NAVY}")
    draw.text((80, 92), "Giả định 10 phút/người/ngày, 30 ngày; chưa gồm text AI, TTS và nhân sự", font=small_font, fill=f"#{MUTED}")

    cf = [18.65, 79.85, 156.35]
    gpu = [497, 497, 497]
    users = [100, 500, 1000]
    max_y = 550
    left, top, right, bottom = 120, 150, 1330, 610
    chart_h = bottom - top
    chart_w = right - left

    for value in range(0, 551, 100):
        y = bottom - (value / max_y) * chart_h
        draw.line((left, y, right, y), fill="#D8DEE8", width=2)
        label = f"${value}"
        bbox = draw.textbbox((0, 0), label, font=small_font)
        draw.text((left - 18 - (bbox[2] - bbox[0]), y - 13), label, font=small_font, fill=f"#{MUTED}")

    group_width = chart_w / 3
    bar_width = 90
    for idx, user_count in enumerate(users):
        center = left + group_width * (idx + 0.5)
        for offset, value, color in ((-55, cf[idx], BLUE), (55, gpu[idx], NAVY)):
            x0 = center + offset - bar_width / 2
            x1 = center + offset + bar_width / 2
            y0 = bottom - (value / max_y) * chart_h
            draw.rounded_rectangle((x0, y0, x1, bottom), radius=8, fill=f"#{color}")
            label = f"${value:.0f}"
            bbox = draw.textbbox((0, 0), label, font=bold_font)
            draw.text(((x0 + x1 - (bbox[2] - bbox[0])) / 2, y0 - 34), label, font=bold_font, fill=f"#{color}")
        user_text = f"{user_count:,}".replace(",", ".") + " người"
        bbox = draw.textbbox((0, 0), user_text, font=label_font)
        draw.text((center - (bbox[2] - bbox[0]) / 2, bottom + 18), user_text, font=label_font, fill=f"#{INK}")

    legend_y = 675
    draw.rounded_rectangle((460, legend_y - 11, 490, legend_y + 19), radius=4, fill=f"#{BLUE}")
    draw.text((505, legend_y - 14), "Cloudflare", font=small_font, fill=f"#{INK}")
    draw.rounded_rectangle((730, legend_y - 11, 760, legend_y + 19), radius=4, fill=f"#{NAVY}")
    draw.text((775, legend_y - 14), "GPU cloud 24/7", font=small_font, fill=f"#{INK}")
    img.save(path, quality=95)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    specs = {
        "Heading 1": (16, BLUE, 12, 6),
        "Heading 2": (13, BLUE, 10, 5),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)


def configure_sections(doc):
    doc.settings.odd_and_even_pages_header_footer = False
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = False

        header = section.header
        hp = header.paragraphs[0]
        hp.paragraph_format.space_after = Pt(0)
        hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        r1 = hp.add_run("BÁO CÁO QUYẾT ĐỊNH CÔNG NGHỆ")
        set_run_font(r1, size=8.5, color=MUTED, bold=True)
        r2 = hp.add_run("\t21/08/2026")
        set_run_font(r2, size=8.5, color=MUTED)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)
        add_page_number(fp)


def set_picture_alt_text(inline_shape, title, description):
    docpr = inline_shape._inline.docPr
    docpr.set("title", title)
    docpr.set("descr", description)


def add_source(doc, title, url, note=None):
    p = doc.add_paragraph()
    apply_num(p, BULLET_NUM_ID)
    p.paragraph_format.space_after = Pt(5)
    add_hyperlink(p, title, url)
    if note:
        r = p.add_run(f" - {note}")
        set_run_font(r, size=10, color=MUTED)


def build_report():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    create_chart(CHART_PATH)

    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    global BULLET_NUM_ID, NUMBER_NUM_ID
    BULLET_NUM_ID, NUMBER_NUM_ID = add_custom_numbering(doc)

    props = doc.core_properties
    props.title = "So sánh Cloudflare Workers AI và model AI tự vận hành"
    props.subject = "Độ khó, chi phí và ngưỡng chuyển đổi với 10 phút audio/người/ngày"
    props.author = "Speaking AI"
    props.last_modified_by = "Speaking AI"
    props.keywords = "Cloudflare Workers AI, Whisper, self-host, ASR, chi phí"
    props.created = datetime(2026, 8, 21, tzinfo=timezone.utc)

    # Page 1: Decision memo masthead.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("ĐÁNH GIÁ CHIẾN LƯỢC ASR")
    set_run_font(r, size=10.5, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Cloudflare Workers AI\nhay model AI tự vận hành?")
    set_run_font(r, size=25, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("So sánh độ khó, chi phí và thời điểm chuyển giải pháp")
    set_run_font(r, size=14, color=MUTED)

    metadata = [
        ("Phạm vi", "100, 500 và 1.000 người dùng; mỗi người nói 10 phút/ngày"),
        ("Mô hình so sánh", "Cloudflare Whisper Large V3 Turbo và self-host/fine-tune model mở tương đương"),
        ("Ngày đánh giá", "21/08/2026"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=10.5, color=NAVY, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=INK)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_callout(
        doc,
        "KHUYẾN NGHỊ",
        "Giữ Cloudflare làm production primary đến ít nhất 1.000 người dùng. Bắt đầu POC model riêng ở chế độ shadow từ 500-1.000 người dùng nếu chất lượng giọng trẻ em là vấn đề. Chỉ cân nhắc chuyển primary từ khoảng 3.000 người dùng trở lên sau khi vượt qua các tiêu chí về chi phí, chất lượng, độ trễ và dự phòng.",
        fill=LIGHT_BLUE,
        accent=NAVY,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_bullet(doc, "Ở 100 người dùng, Cloudflare khoảng 17,00-20,30 USD/tháng.", BULLET_NUM_ID, "Ở 100 người dùng")
    add_bullet(doc, "Ở 1.000 người dùng, Cloudflare khoảng 154,70-158,00 USD/tháng.", BULLET_NUM_ID, "Ở 1.000 người dùng")
    add_bullet(doc, "Một GPU cloud luôn chạy có giá khoảng 497-929 USD/tháng trước chi phí nhân sự và dự phòng.", BULLET_NUM_ID, "Một GPU cloud luôn chạy")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    inline = p.add_run().add_picture(str(CHART_PATH), width=Inches(5.75))
    set_picture_alt_text(
        inline,
        "Biểu đồ chi phí ASR",
        "So sánh chi phí Cloudflare và một GPU cloud chạy liên tục ở 100, 500 và 1.000 người dùng.",
    )
    cap = doc.add_paragraph("Hình 1. Cloudflare vẫn rẻ hơn GPU cloud luôn chạy ở quy mô đến 1.000 người dùng.", style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(0)

    # Page 2: Assumptions and Cloudflare costs.
    add_page_break(doc)
    add_heading(doc, "1. Phạm vi và giả định", 1)
    add_body(
        doc,
        "Báo cáo lấy 30 ngày/tháng và 10 phút audio/người/ngày. Bảng chi phí chính chỉ tính ASR, vì số phút nói không xác định được số token dịch văn bản hoặc số ký tự TTS. Chi phí text AI và TTS được trình bày riêng ở cuối báo cáo.",
    )
    add_bullet(doc, "Cloudflare sử dụng @cf/openai/whisper-large-v3-turbo với giá 0,00051 USD/phút audio.", BULLET_NUM_ID)
    add_bullet(doc, "Workers Paid tối thiểu 5 USD/tháng; quota AI miễn phí có giá trị tối đa khoảng 3,30 USD/tháng và dùng chung cho các model.", BULLET_NUM_ID)
    add_bullet(doc, "Tỷ giá quy đổi minh họa: 26.000 VND/USD; chưa gồm thuế.", BULLET_NUM_ID)
    add_bullet(doc, "Model tự build được hiểu là self-host hoặc fine-tune model có sẵn, không phải huấn luyện foundation model từ đầu.", BULLET_NUM_ID)

    add_heading(doc, "Khối lượng sử dụng", 2)
    add_table(
        doc,
        ["Người dùng", "Audio/ngày", "Audio/tháng", "Giờ audio/tháng"],
        [
            ["100", "1.000 phút", "30.000 phút", "500 giờ"],
            ["500", "5.000 phút", "150.000 phút", "2.500 giờ"],
            ["1.000", "10.000 phút", "300.000 phút", "5.000 giờ"],
        ],
        [1500, 2400, 2760, 2700],
        [WD_ALIGN_PARAGRAPH.CENTER] * 4,
        font_size=9.6,
    )

    add_heading(doc, "2. Chi phí Cloudflare Workers AI", 1)
    add_callout(
        doc,
        "CÔNG THỨC",
        "Chi phí ASR = tổng số phút audio x 0,00051 USD. Tổng tháng = ASR + 5 USD Workers Paid - quota miễn phí còn lại.",
        fill=LIGHT_GRAY,
        accent=BLUE,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_table(
        doc,
        ["Người dùng", "ASR", "Tổng Cloudflare", "Quy đổi VND", "Nhận xét"],
        [
            ["100", "$15,30", "$17,00-$20,30", "442.000-528.000đ", "Rất thấp"],
            ["500", "$76,50", "$78,20-$81,50", "2,03-2,12 triệu", "Phù hợp"],
            ["1.000", "$153,00", "$154,70-$158,00", "4,02-4,11 triệu", "Vẫn rẻ hơn GPU riêng"],
        ],
        [1200, 1450, 2000, 2260, 2450],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9.1,
        header_fill=LIGHT_BLUE,
    )
    p = doc.add_paragraph(style="Caption")
    p.add_run("Nguồn: ")
    add_hyperlink(p, "Cloudflare Whisper Large V3 Turbo", "https://developers.cloudflare.com/workers-ai/models/whisper-large-v3-turbo/")
    p.add_run("; ")
    add_hyperlink(p, "Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/")
    p.add_run("; ")
    add_hyperlink(p, "Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/")

    # Page 3: self-host costs.
    add_page_break(doc)
    add_heading(doc, "3. Chi phí model AI tự vận hành", 1)
    add_body(
        doc,
        "Phương án thực tế là triển khai Whisper Large V3 Turbo hoặc một model tiếng Việt tương đương bằng faster-whisper, sau đó fine-tune nếu dữ liệu giọng trẻ em đủ chất lượng. Huấn luyện foundation model từ đầu không phù hợp với quy mô 100-1.000 người dùng.",
    )

    add_heading(doc, "3.1. GPU cloud luôn chạy", 2)
    add_body(
        doc,
        "Ứng dụng hội thoại cần phản hồi nhanh, vì vậy GPU thường phải luôn sẵn sàng. Giá tham khảo của Lambda là 0,69 USD/giờ cho Quadro RTX 6000 và 1,29 USD/giờ cho A10, tương đương khoảng 497-929 USD/tháng cho một GPU chạy 24/7.",
    )
    add_table(
        doc,
        ["Người dùng", "Một GPU/tháng", "Hai GPU HA/tháng", "Đánh giá"],
        [
            ["100", "$497-$929", "$994-$1.858", "GPU gần như nhàn rỗi"],
            ["500", "$497-$929", "$994-$1.858", "Đắt hơn Cloudflare rõ rệt"],
            ["1.000", "$497-$929", "$994-$1.858", "Có thể đủ tải trung bình; phải kiểm tra tải đỉnh"],
        ],
        [1450, 2050, 2250, 3610],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9.2,
    )

    add_heading(doc, "3.2. GPU bật theo nhu cầu", 2)
    add_body(
        doc,
        "Nếu có thể scale-to-zero và chấp nhận cold start, chi phí compute thuần có thể thấp hơn. Ước tính dưới đây giả định tốc độ 10-30 lần thời gian thực trên GPU 0,69 USD/giờ. Đây chưa phải giá production vì chưa gồm thời gian chờ request, tải model, CPU, storage, endpoint và dự phòng.",
    )
    add_table(
        doc,
        ["Người dùng", "GPU compute thuần", "So với Cloudflare"],
        [
            ["100", "$11,50-$34,50", "Có thể gần Cloudflare, nhưng cold start không phù hợp hội thoại"],
            ["500", "$57,50-$172,50", "Chỉ cạnh tranh khi batching và sử dụng GPU tốt"],
            ["1.000", "$115-$345", "Biên độ lớn; cần benchmark bằng audio thật"],
        ],
        [1500, 2500, 5360],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9.3,
        header_fill=PALE_GOLD,
    )

    add_heading(doc, "3.3. Máy GPU tự mua", 2)
    add_bullet(doc, "Khấu hao phần cứng: khoảng 100-180 USD/tháng cho một máy, tùy cấu hình và thời gian khấu hao.", BULLET_NUM_ID)
    add_bullet(doc, "Điện, mạng, storage và backup: khoảng 50-120 USD/tháng.", BULLET_NUM_ID)
    add_bullet(doc, "Tổng một máy: khoảng 150-300 USD/tháng; production có dự phòng thường gấp đôi.", BULLET_NUM_ID)
    add_bullet(doc, "Chưa tính chi phí kỹ sư ML/DevOps, giám sát, cập nhật CUDA, bảo mật và xử lý sự cố.", BULLET_NUM_ID)
    p = doc.add_paragraph(style="Caption")
    p.add_run("Nguồn GPU cloud: ")
    add_hyperlink(p, "Lambda GPU instances", "https://lambda.ai/instances")
    p.add_run(". Hiệu năng phải được kiểm tra lại trên audio ngắn của dự án; xem ")
    add_hyperlink(p, "Faster Whisper benchmark", "https://github.com/SYSTRAN/faster-whisper")
    p.add_run(".")

    # Page 4: difficulty.
    add_page_break(doc)
    add_heading(doc, "4. So sánh độ khó", 1)
    add_table(
        doc,
        ["Tiêu chí", "Cloudflare", "Self-host model", "Fine-tune model riêng"],
        [
            ["Độ khó tổng thể", "2/10", "6/10", "8/10"],
            ["Thời gian POC", "Đã tích hợp", "2-4 tuần", "1-3 tháng"],
            ["Quản lý GPU/CUDA", "Không", "Bắt buộc", "Bắt buộc"],
            ["Autoscaling", "Cloudflare xử lý", "Tự xây dựng", "Tự xây dựng"],
            ["Monitoring và retry", "Đơn giản", "Tự triển khai", "Tự triển khai"],
            ["Chất lượng giọng trẻ em", "Model tổng quát", "Gần model gốc", "Có thể tốt hơn nếu dữ liệu tốt"],
            ["Quyền kiểm soát dữ liệu", "Thấp hơn", "Cao", "Cao"],
            ["Rủi ro downtime", "Thấp", "Trung bình-cao", "Trung bình-cao"],
            ["Chi phí nhân sự", "Thấp", "Trung bình-cao", "Cao"],
        ],
        [2380, 1900, 2400, 2680],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
        header_fill=LIGHT_BLUE,
    )

    add_heading(doc, "Phần khó nhất của model riêng", 2)
    add_bullet(doc, "Thu thập audio có consent và gắn nhãn chính xác theo vùng miền, độ tuổi và môi trường.", BULLET_NUM_ID)
    add_bullet(doc, "Đánh giá Word Error Rate và độ đúng ngữ nghĩa trên một golden set ổn định.", BULLET_NUM_ID)
    add_bullet(doc, "Tối ưu batching nhưng vẫn giữ độ trễ thấp cho các đoạn hội thoại ngắn.", BULLET_NUM_ID)
    add_bullet(doc, "Autoscaling GPU khi nhiều trẻ nói đồng thời và xử lý tải đỉnh gấp đôi bình thường.", BULLET_NUM_ID)
    add_bullet(doc, "Thiết lập dự phòng, rollback, bảo mật dữ liệu, theo dõi drift và tái huấn luyện.", BULLET_NUM_ID)

    add_callout(
        doc,
        "LƯU Ý",
        "10.000 audio không tự động tạo ra model tốt. Giá trị phụ thuộc vào tổng số giờ, độ đa dạng, độ sạch của nhãn, quyền sử dụng dữ liệu và mức độ đại diện cho giọng trẻ em thực tế.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    # Page 5: break-even and stop conditions.
    add_page_break(doc)
    add_heading(doc, "5. Điểm hòa vốn và thời điểm dừng", 1)
    add_callout(
        doc,
        "CHI PHÍ BIẾN ĐỔI",
        "Một người dùng tạo 300 phút audio/tháng. Chi phí ASR Cloudflare tương ứng khoảng 0,153 USD/người/tháng.",
        fill=LIGHT_GRAY,
        accent=BLUE,
    )

    add_heading(doc, "Ngưỡng quyết định theo quy mô", 2)
    add_body(doc, "100-1.000 người dùng: giữ Cloudflare production primary.", bold_lead="100-1.000 người dùng:")
    add_body(doc, "500-1.000 người dùng: có thể bắt đầu POC model riêng ở chế độ shadow nếu mục tiêu là cải thiện chất lượng, không phải chỉ để giảm chi phí.", bold_lead="500-1.000 người dùng:")
    add_body(doc, "1.300-2.600 người dùng: điểm hòa vốn lý thuyết của một đến hai máy GPU tự sở hữu, chưa tính nhân sự.", bold_lead="1.300-2.600 người dùng:")
    add_body(doc, "Khoảng 3.200 người dùng: một GPU cloud 497 USD/tháng bắt đầu gần hòa vốn với Cloudflare ASR.", bold_lead="Khoảng 3.200 người dùng:")
    add_body(doc, "5.000-10.000 người dùng: vùng cân nhắc chuyển thực tế sau khi cộng kỹ sư, HA, downtime và chi phí vận hành.", bold_lead="5.000-10.000 người dùng:")

    add_heading(doc, "Điều kiện để dừng Cloudflare làm primary", 2)
    conditions = [
        ("Chi phí đầy đủ", "dưới 0,0004 USD/phút và thấp hơn Cloudflare ít nhất 20%."),
        ("Chất lượng", "bằng hoặc tốt hơn Cloudflare trên bộ audio chuẩn của dự án."),
        ("Mức cải thiện", "WER hoặc tỷ lệ hiểu sai giọng trẻ em tốt hơn ít nhất 10-15%."),
        ("Độ trễ", "p95 đáp ứng mục tiêu trải nghiệm hội thoại của dự án."),
        ("Khả năng chịu tải", "xử lý được ít nhất gấp đôi tải đỉnh dự kiến."),
        ("Độ ổn định", "chạy shadow/canary tối thiểu 30 ngày không có lỗi nghiêm trọng."),
        ("Fallback", "có GPU dự phòng và quy trình rollback rõ ràng."),
    ]
    for label, detail in conditions:
        add_numbered(doc, f"{label}: {detail}", NUMBER_NUM_ID, f"{label}:")

    add_heading(doc, "Khi nào dừng hoặc hoãn self-build", 2)
    add_bullet(doc, "Chất lượng sau POC chỉ tương đương Cloudflare hoặc cải thiện dưới 10%.", BULLET_NUM_ID)
    add_bullet(doc, "Chi phí đầy đủ vẫn cao hơn 0,00051 USD/phút.", BULLET_NUM_ID)
    add_bullet(doc, "GPU sử dụng dưới 30% hoặc không thể giữ độ trễ ổn định vào giờ cao điểm.", BULLET_NUM_ID)
    add_bullet(doc, "Phải dành hơn một ngày kỹ thuật mỗi tuần chỉ để vận hành model.", BULLET_NUM_ID)
    add_bullet(doc, "Dữ liệu chưa đủ sạch hoặc chưa có quyền phù hợp để fine-tune.", BULLET_NUM_ID)

    # Page 6: roadmap and caveats.
    add_page_break(doc)
    add_heading(doc, "6. Lộ trình đề xuất", 1)
    add_table(
        doc,
        ["Giai đoạn", "Quy mô", "Hành động", "Tiêu chí ra quyết định"],
        [
            ["Hiện tại", "100-500", "Cloudflare primary; đo phút audio, latency, lỗi và chi phí", "Ổn định production"],
            ["Chuẩn bị", "500-1.000", "Tạo golden set 300-500 audio; benchmark model riêng ở shadow", "Chất lượng có cải thiện rõ"],
            ["Đánh giá", "2.000-3.000", "POC self-host; đo cost/min, peak load và nhân sự vận hành", "TCO thấp hơn ít nhất 20%"],
            ["Chuyển dần", ">3.000", "Canary 5%-20%-50%; Cloudflare vẫn là fallback", "30 ngày ổn định"],
            ["Tối ưu", "5.000+", "Xem xét self-host primary hoặc mô hình hybrid", "Đạt SLA và mục tiêu chất lượng"],
        ],
        [1350, 1450, 3820, 2740],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.8,
        header_fill=PALE_GREEN,
    )

    add_heading(doc, "7. Chi phí ngoài ASR", 1)
    add_body(
        doc,
        "Ứng dụng hiện tại còn dùng text AI và TTS. Hai khoản này không thể suy ra chính xác chỉ từ 10 phút audio/ngày vì phụ thuộc số lượt hội thoại, số token và số câu tiếng Anh mới được tạo.",
    )
    add_bullet(doc, "Llama 4 Scout: 0,27 USD/triệu input token và 0,85 USD/triệu output token.", BULLET_NUM_ID)
    add_bullet(doc, "Deepgram Aura TTS: 0,015 USD/1.000 ký tự. Cache audio giúp lượt phát lại không tạo thêm chi phí TTS.", BULLET_NUM_ID)
    add_bullet(doc, "Nếu mỗi câu trả lời đều mới và không cache, TTS có thể lớn hơn chi phí ASR; cần đo số ký tự TTS duy nhất/tháng.", BULLET_NUM_ID)

    add_callout(
        doc,
        "KẾT LUẬN",
        "Ở quy mô 100-1.000 người dùng, Cloudflare là lựa chọn hợp lý nhất về tổng chi phí và độ khó. Model riêng chỉ nên được đầu tư sớm khi có yêu cầu mạnh về chất lượng giọng trẻ em, quyền kiểm soát dữ liệu hoặc năng lực ML nội bộ; không nên chuyển chỉ vì kỳ vọng tiết kiệm hạ tầng.",
        fill=PALE_GREEN,
        accent=GREEN,
    )

    add_heading(doc, "8. Nguồn tham khảo", 1)
    add_source(doc, "Cloudflare Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/", "giá request và CPU time")
    add_source(doc, "Cloudflare Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/", "quota Neurons và phí vượt mức")
    add_source(doc, "Whisper Large V3 Turbo pricing", "https://developers.cloudflare.com/workers-ai/models/whisper-large-v3-turbo/", "0,00051 USD/phút audio")
    add_source(doc, "Llama 4 Scout pricing", "https://developers.cloudflare.com/workers-ai/models/llama-4-scout-17b-16e-instruct/", "giá token")
    add_source(doc, "Deepgram Aura pricing", "https://developers.cloudflare.com/workers-ai/models/aura-1/", "giá TTS theo ký tự")
    add_source(doc, "Lambda GPU instances", "https://lambda.ai/instances", "giá GPU cloud tham khảo")
    add_source(doc, "SYSTRAN faster-whisper", "https://github.com/SYSTRAN/faster-whisper", "triển khai và benchmark inference")
    add_source(doc, "OpenAI Whisper Large V3 Turbo model card", "https://huggingface.co/openai/whisper-large-v3-turbo", "kiến trúc và phạm vi model")

    # Structural settings that Word/LibreOffice both honor.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_report()

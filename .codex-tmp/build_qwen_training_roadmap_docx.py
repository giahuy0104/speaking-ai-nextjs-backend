# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\Code\HuaMei\App_noi\be\9_10th8\speaking-ai-nextjs-backend")
OUTPUT = ROOT / "deliverables" / "Lo_trinh_hoc_va_trien_khai_AI_ngon_ngu_Qwen.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
INK = "202124"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "A56A00"
WHITE = "FFFFFF"
BORDER = "C9D2DC"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_SIDE_DXA = 120


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    if run._element.get_or_add_rPr().rFonts is None:
        rfonts = OxmlElement("w:rFonts")
        run._element.get_or_add_rPr().insert(0, rfonts)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=CELL_TOP_BOTTOM_DXA, start=CELL_SIDE_DXA,
                     bottom=CELL_TOP_BOTTOM_DXA, end=CELL_SIDE_DXA):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    assert sum(widths_dxa) == CONTENT_DXA, (widths_dxa, sum(widths_dxa))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    set_table_borders(table)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border_left(paragraph, color=BLUE, size=16, space=10):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("Trang ")
    set_run_font(prefix, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    code = styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor.from_string("273142")
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.08

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)


def add_numbering_definitions(doc):
    numbering = doc.part.numbering_part.element

    def add_abstract(abstract_id, num_fmt, text, left=540, hanging=270, font="Calibri"):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        txt = OxmlElement("w:lvlText")
        txt.set(qn("w:val"), text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), font)
        r_fonts.set(qn("w:hAnsi"), font)
        r_pr.append(r_fonts)
        lvl.extend([start, fmt, txt, jc, p_pr, r_pr])
        abstract.append(lvl)
        numbering.append(abstract)

    def add_num(num_id, abstract_id):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    add_abstract(100, "bullet", "•")
    add_num(100, 100)
    add_abstract(101, "decimal", "%1.")
    add_num(101, 101)
    add_abstract(102, "bullet", "☐", font="Segoe UI Symbol")
    add_num(102, 102)


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def set_manual_list_layout(paragraph):
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_manual_list_layout(p)
    marker = p.add_run("•\t")
    set_run_font(marker, bold=True, color=BLUE)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_numbered(doc, text, index):
    p = doc.add_paragraph()
    set_manual_list_layout(p)
    marker = p.add_run(f"{index}.\t")
    set_run_font(marker, bold=True, color=BLUE)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_check(doc, text):
    p = doc.add_paragraph()
    set_manual_list_layout(p)
    marker = p.add_run("☐\t")
    set_run_font(marker, name="Segoe UI Symbol", color=INK)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_callout(doc, label, text, fill=CALLOUT, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.18
    set_paragraph_shading(p, fill)
    set_paragraph_border_left(p, accent)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, color=INK)
    return p


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    set_paragraph_shading(p, LIGHT_GRAY)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=9, color="273142")
    return p


def add_table(doc, headers, rows, widths_dxa, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=True, color=DARK_BLUE)
    set_repeat_table_header(table.rows[0])

    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            if row_idx % 2 == 1:
                set_cell_shading(cells[idx], "FAFBFC")
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_section_intro(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text)
    set_run_font(r, size=11, color=MUTED, italic=True)
    return p


def add_page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_styles(doc)
    add_numbering_definitions(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("Cẩm nang triển khai AI ngôn ngữ | Qwen cho ứng dụng Việt-Anh")
    set_run_font(hr, size=8.5, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)

    # Cover: editorial_cover pattern with a named title override.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    kr = kicker.add_run("CẨM NANG THỰC HÀNH DÀNH CHO NGƯỜI MỚI")
    set_run_font(kr, size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    tr = title.add_run("LỘ TRÌNH TỰ HỌC VÀ TRIỂN KHAI\nAI NGÔN NGỮ")
    set_run_font(tr, size=28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    sr = subtitle.add_run("Qwen3-4B / Qwen3.5-4B cho ứng dụng luyện nói Việt-Anh")
    set_run_font(sr, size=14, color=DARK_BLUE)

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.paragraph_format.space_after = Pt(48)
    dr = desc.add_run("Từ kiến thức nền tảng đến dữ liệu, QLoRA, tích hợp Next.js và vận hành production")
    set_run_font(dr, size=10.5, italic=True, color=MUTED)

    metrics = add_table(
        doc,
        ["Mục tiêu đầu tiên", "Thời gian tham khảo", "Hạ tầng pilot"],
        [["Self-host và benchmark", "6-8 tuần học bán thời gian", "01 GPU NVIDIA 24 GB"]],
        [3120, 3120, 3120],
        font_size=9.5,
    )
    metrics.alignment = WD_TABLE_ALIGNMENT.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(36)
    meta.paragraph_format.space_after = Pt(0)
    mr = meta.add_run("Phiên bản 1.0 | Tháng 8 năm 2026")
    set_run_font(mr, size=10, bold=True, color=NAVY)

    add_page_break(doc)

    # Executive orientation.
    doc.add_heading("1. Định hướng đúng ngay từ đầu", level=1)
    add_section_intro(doc, "Mục tiêu của tài liệu là giúp một người mới đi từ sử dụng model có sẵn đến fine-tune và tích hợp an toàn, không đi vòng qua bài toán pretraining quy mô lớn.")
    add_callout(
        doc,
        "Khuyến nghị cốt lõi",
        "Không train ngay. Hãy chạy Qwen nguyên bản, tạo bộ test chuẩn, benchmark với Cloudflare và chỉ fine-tune khi đã xác định được nhóm lỗi lặp lại.",
    )
    add_bullet(doc, "Giai đoạn 1 chỉ thay LLM dịch văn bản; tiếp tục giữ Cloudflare cho ASR và TTS.")
    add_bullet(doc, "Không train Qwen từ đầu; sử dụng Supervised Fine-Tuning (SFT) với QLoRA.")
    add_bullet(doc, "Dữ liệu chất lượng và quy trình đánh giá quan trọng hơn số lượng mẫu.")
    add_bullet(doc, "Cloudflare cần được giữ làm fallback trong thời gian shadow và canary.")

    doc.add_heading("Kết quả thực tế có thể kỳ vọng", level=2)
    add_table(
        doc,
        ["Mức độ", "Thời gian tham khảo", "Kết quả"],
        [
            ["Pilot kỹ thuật", "2-4 tuần nếu đã biết Python/Docker", "Qwen chạy qua API, có benchmark ban đầu"],
            ["Người biết lập trình, mới với AI", "6-8 tuần", "Tự chạy QLoRA và đánh giá adapter"],
            ["Production có kiểm soát", "10-12 tuần", "Feature flag, fallback, monitoring, rollback"],
            ["Train LLM từ đầu", "Không phù hợp cá nhân", "Cần đội ngũ, dữ liệu và cụm GPU lớn"],
        ],
        [1700, 2600, 5060],
    )

    doc.add_heading("Bản đồ nội dung", level=2)
    for index, item in enumerate([
        "Hiểu model, tham số và chi phí hạ tầng",
        "Xác định kiến thức cần học và những gì có thể bỏ qua",
        "Xây kho dữ liệu Việt-Anh có kiểm duyệt",
        "Chạy Qwen self-host và benchmark trước khi train",
        "Fine-tune bằng SFT + QLoRA",
        "Tích hợp vào backend Next.js và triển khai production",
    ], start=1):
        add_numbered(doc, item, index)

    # Fundamentals.
    doc.add_heading("2. Hiểu đúng về model và huấn luyện", level=1)
    doc.add_heading("2.1. Ba khái niệm không được nhầm lẫn", level=2)
    add_table(
        doc,
        ["Khái niệm", "Ý nghĩa", "Áp dụng cho dự án"],
        [
            ["Inference", "Dùng trọng số có sẵn để sinh kết quả", "Chạy Qwen trên server riêng"],
            ["Fine-tuning", "Điều chỉnh hành vi model bằng dữ liệu miền", "Dạy cách dịch câu trẻ em trung thành hơn"],
            ["Pretraining", "Huấn luyện nền tảng từ lượng dữ liệu khổng lồ", "Không cần và không phù hợp"],
        ],
        [1700, 3400, 4260],
    )

    doc.add_heading("2.2. Bốn tỷ tham số nghĩa là gì?", level=2)
    p = doc.add_paragraph()
    p.add_run("Qwen3-4B chứa khoảng ").bold = False
    r = p.add_run("4.000.000.000 giá trị số có thể học được")
    set_run_font(r, bold=True, color=DARK_BLUE)
    p.add_run(". Chúng nằm trong embedding, attention, feed-forward và các lớp đầu ra. Đây không phải bốn tỷ từ hay bốn tỷ câu trả lời; kiến thức được phân tán trên nhiều tham số.")
    add_code(doc, "Token tiếng Việt → Embedding → 36 lớp xử lý → Xác suất token tiếng Anh → Câu dịch")
    add_table(
        doc,
        ["Định dạng", "Số byte/tham số", "Dung lượng thô cho 4B"],
        [
            ["FP32", "4 byte", "Khoảng 16 GB"],
            ["BF16 / FP16", "2 byte", "Khoảng 8 GB"],
            ["INT8", "1 byte", "Khoảng 4 GB"],
            ["INT4", "0,5 byte", "Khoảng 2 GB"],
        ],
        [2600, 2600, 4160],
    )
    add_callout(doc, "Lưu ý", "Quantize INT4 không làm model giảm còn một tỷ tham số. Model vẫn có bốn tỷ tham số, nhưng mỗi giá trị được lưu bằng ít bit hơn.")

    doc.add_heading("2.3. Vì sao training tốn hạ tầng?", level=2)
    add_bullet(doc, "Inference chỉ chạy forward; training còn phải lưu activation, tính loss và chạy backpropagation.")
    add_bullet(doc, "Full training phải giữ trọng số, gradient, bản FP32, optimizer states và bộ nhớ trung gian.")
    add_bullet(doc, "Sequence length và batch size càng lớn thì VRAM càng tăng.")
    add_bullet(doc, "Mỗi mẫu được xử lý nhiều lần qua các epoch.")
    add_table(
        doc,
        ["Phương pháp", "VRAM thực hành ước tính", "Khuyến nghị"],
        [
            ["Full fine-tuning 4B", "64-100+ GB", "Không dùng cho pilot"],
            ["LoRA BF16", "20-40 GB", "Có thể dùng nếu có GPU lớn"],
            ["QLoRA 4-bit", "12-24 GB", "Phù hợp nhất"],
            ["Inference INT4", "4-8 GB", "Dùng để thử nghiệm"],
        ],
        [2600, 2800, 3960],
    )

    # Learning map.
    doc.add_heading("3. Kiến thức cần học", level=1)
    add_section_intro(doc, "Mục tiêu là đủ năng lực xây một vòng đời AI hoàn chỉnh, không phải học mọi chủ đề trong Machine Learning.")

    doc.add_heading("3.1. Năng lực kỹ thuật tối thiểu", level=2)
    for text in [
        "Python: hàm, class, đọc JSON/CSV, virtual environment và pip.",
        "Linux command line: file, process, biến môi trường và quyền truy cập.",
        "Git: branch, commit, tag và quản lý phiên bản cấu hình/dataset.",
        "Docker: image, container, volume, port và GPU runtime.",
        "HTTP/REST: JSON, status code, timeout, authentication và retry.",
        "GPU cơ bản: CUDA, VRAM, precision, quantization và concurrency.",
        "Next.js/TypeScript: xây provider adapter và feature flag.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("3.2. Khái niệm AI bắt buộc", level=2)
    add_table(
        doc,
        ["Khái niệm", "Điều cần hiểu"],
        [
            ["Token / tokenizer", "Model đọc token số, không đọc trực tiếp chữ"],
            ["Context length", "Giới hạn token đầu vào và đầu ra"],
            ["Loss / epoch / batch", "Cách đo lỗi và tổ chức vòng lặp train"],
            ["Learning rate", "Mức thay đổi trọng số ở mỗi bước"],
            ["Overfitting", "Nhớ train set nhưng làm kém với câu mới"],
            ["Sampling", "Temperature, top-p và tính ổn định đầu ra"],
            ["LoRA / QLoRA", "Fine-tune adapter nhỏ thay vì toàn bộ model"],
            ["Hallucination", "Model thêm hoặc đổi thông tin không có trong nguồn"],
        ],
        [2500, 6860],
    )

    doc.add_heading("3.3. Chưa cần học ở giai đoạn đầu", level=2)
    for text in [
        "Viết CUDA kernel hoặc tối ưu compiler cho GPU.",
        "Pretraining Transformer từ đầu và huấn luyện tokenizer mới.",
        "Distributed training nhiều node, NVLink hoặc InfiniBand.",
        "RLHF, reward model, DPO hoặc GRPO.",
        "Voice cloning và tự train TTS khi chưa ổn định phần LLM.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("3.4. Chọn đúng kỹ thuật", level=2)
    add_table(
        doc,
        ["Nhu cầu", "Giải pháp phù hợp"],
        [
            ["Model cần tra cứu tài liệu mới", "RAG"],
            ["Model dịch sai phong cách hoặc quy tắc", "Prompt, rule hoặc SFT/LoRA"],
            ["Nhận dạng giọng nói sai", "Tối ưu/fine-tune ASR"],
            ["Muốn giọng đọc riêng", "Fine-tune TTS"],
            ["Muốn dữ liệu không rời hạ tầng", "Self-host inference"],
        ],
        [4200, 5160],
    )

    # Architecture.
    doc.add_heading("4. Kiến trúc cho ứng dụng hiện tại", level=1)
    doc.add_heading("4.1. Cloudflare đang đảm nhiệm nhiều vai trò", level=2)
    add_table(
        doc,
        ["Công đoạn", "Hiện tại", "Ứng viên self-host"],
        [
            ["ASR tiếng Việt", "Whisper Large V3 Turbo", "faster-whisper large-v3-turbo"],
            ["Việt → Anh", "Llama 4 Scout", "Qwen3-4B-Instruct-2507 hoặc Qwen3.5-4B"],
            ["TTS tiếng Anh", "Deepgram Aura-1", "Kokoro-82M"],
            ["Cache audio", "Local/Postgres/R2 tùy cấu hình", "Giữ local/Postgres hoặc S3/MinIO"],
        ],
        [2300, 3300, 3760],
    )

    doc.add_heading("4.2. Phạm vi triển khai đầu tiên", level=2)
    add_code(doc, "Ứng dụng → ASR hiện tại → Rule/Text cache → Qwen self-host → TTS hiện tại")
    add_callout(doc, "Quyết định kiến trúc", "Chỉ thay bước dịch văn bản trong giai đoạn đầu. ASR và TTS là hai bài toán dữ liệu/model riêng, không nên thay đồng thời.")

    doc.add_heading("4.3. Dịch vụ model độc lập", level=2)
    add_code(doc, "Next.js backend\n   └─ HTTPS /v1/chat/completions\n        └─ vLLM trên GPU server\n             └─ Qwen3-4B-Instruct-2507")
    add_bullet(doc, "Không nạp model trực tiếp vào tiến trình Next.js/Vercel.")
    add_bullet(doc, "Đặt model trên GPU service riêng, có HTTPS và API key.")
    add_bullet(doc, "Giới hạn context ở 4K vì prompt và output của ứng dụng ngắn.")
    add_bullet(doc, "Giữ timeout, circuit breaker và fallback Cloudflare.")

    doc.add_heading("4.4. Cấu hình khởi đầu", level=2)
    add_code(doc, "TEXT_AI_PROVIDER=self_hosted\nSELF_HOSTED_LLM_BASE_URL=https://gpu.example.com/v1\nSELF_HOSTED_LLM_API_KEY=<secret>\nSELF_HOSTED_LLM_MODEL=Qwen/Qwen3-4B-Instruct-2507\nSELF_HOSTED_LLM_TIMEOUT_MS=2500")

    # Data.
    doc.add_heading("5. Xây kho dữ liệu Việt-Anh", level=1)
    add_section_intro(doc, "Kho ngôn ngữ chỉ trở thành dữ liệu train khi mỗi mẫu có mục tiêu, nhãn chuẩn, metadata và trạng thái kiểm duyệt rõ ràng.")
    doc.add_heading("5.1. Dữ liệu cho ba model khác nhau", level=2)
    add_table(
        doc,
        ["Model", "Đầu vào", "Nhãn chuẩn"],
        [
            ["ASR / Whisper", "File audio trẻ nói", "Transcript tiếng Việt"],
            ["LLM / Qwen", "Câu tiếng Việt", "Bản dịch tiếng Anh"],
            ["TTS / Kokoro", "Văn bản tiếng Anh", "Audio người đọc tương ứng"],
        ],
        [2400, 3400, 3560],
    )
    add_callout(doc, "Quan trọng", "Audio không thể đưa trực tiếp vào Qwen3-4B. Muốn train Qwen cần transcript tiếng Việt và bản dịch tiếng Anh đã được con người xác nhận.")

    doc.add_heading("5.2. Cấu trúc một mẫu dữ liệu", level=2)
    add_code(doc, '{\n  "id": "vi_en_000001",\n  "vietnamese": "Mẹ ơi, con không muốn mua cái này.",\n  "english": "Mom, I don\'t want to buy this.",\n  "context": "outside",\n  "region": "south",\n  "speech_act": "refusal",\n  "review_status": "human_approved",\n  "split": "train"\n}')

    doc.add_heading("5.3. Quy tắc chất lượng", level=2)
    for text in [
        "Mỗi bản dịch phải được con người xác nhận; output Cloudflare chỉ dùng làm bản nháp.",
        "Gắn nhãn phủ định, chủ thể, người nghe, màu sắc, số lượng, thời gian và địa điểm.",
        "Bao phủ câu hỏi, yêu cầu, xin phép, mong muốn, từ chối và câu chưa hoàn chỉnh.",
        "Đưa từ vùng miền Bắc-Trung-Nam vào guideline, tránh thay sai từ đồng âm.",
        "Loại bản ghi trùng hoặc gần trùng trước khi chia tập.",
    ]:
        add_check(doc, text)

    doc.add_heading("5.4. Chia dữ liệu", level=2)
    add_table(
        doc,
        ["Tập", "Tỷ lệ gợi ý", "Quy tắc"],
        [
            ["Train", "80%", "Dùng cập nhật adapter"],
            ["Validation", "10%", "Chọn checkpoint và phát hiện overfitting"],
            ["Test", "10%", "Không được dùng trong train hoặc chọn tham số"],
        ],
        [2300, 2300, 4760],
    )
    add_bullet(doc, "Chia theo cụm ngữ nghĩa để câu chỉ khác màu sắc không bị rơi vào hai tập khác nhau.")
    add_bullet(doc, "Với ASR, chia theo người nói để audio của cùng một trẻ không xuất hiện ở cả train và test.")

    doc.add_heading("5.5. Quy mô khởi đầu", level=2)
    add_table(
        doc,
        ["Số mẫu", "Mục đích"],
        [
            ["500-1.000", "Golden test set và benchmark model gốc"],
            ["2.000-5.000", "Thử QLoRA đầu tiên"],
            ["Khoảng 10.000", "Dataset miền tương đối tốt nếu đã human-approved"],
        ],
        [2700, 6660],
    )

    # Roadmap.
    doc.add_heading("6. Lộ trình học và triển khai trong 8 tuần", level=1)
    add_section_intro(doc, "Giả định: đã biết lập trình cơ bản và dành 1-2 giờ mỗi ngày. Nếu học toàn thời gian, các mốc có thể rút ngắn; nếu chưa biết lập trình, cần thêm 2-4 tháng nền tảng.")
    roadmap_rows = [
        ["1", "Nền tảng LLM", "Token, inference, fine-tuning, loss, epoch, overfitting", "Giải thích được vòng đời model"],
        ["2", "Python & GPU", "Python, Transformers, CUDA/VRAM, môi trường", "Chạy được model từ Hugging Face"],
        ["3", "Self-host", "vLLM, API, authentication, latency", "Có /v1/chat/completions nội bộ"],
        ["4", "Dữ liệu", "Guideline, gán nhãn, dedup, split", "Có 500-1.000 mẫu golden"],
        ["5", "Benchmark", "So Qwen với Cloudflare theo từng nhóm lỗi", "Biết có cần fine-tune hay không"],
        ["6", "QLoRA", "SFT, adapter, checkpoint, training curve", "Có adapter đầu tiên"],
        ["7", "Đánh giá", "Test kín, overfitting, quantization", "Chứng minh model tốt hơn baseline"],
        ["8", "Tích hợp", "Feature flag, shadow, canary, rollback", "Pilot chạy trong ứng dụng"],
    ]
    add_table(
        doc,
        ["Tuần", "Trọng tâm", "Nội dung", "Đầu ra"],
        roadmap_rows,
        [800, 1900, 3900, 2760],
        font_size=8.8,
    )

    doc.add_heading("Mốc hoàn thành pilot", level=2)
    for text in [
        "Model self-host trả được đúng cấu trúc API.",
        "Golden set có reviewer và không rò rỉ dữ liệu cá nhân.",
        "Có báo cáo so sánh Cloudflare, Qwen gốc và Qwen + LoRA.",
        "Ứng dụng có feature flag và fallback.",
        "Có quy trình rollback khi latency hoặc chất lượng giảm.",
    ]:
        add_check(doc, text)

    # Training.
    doc.add_heading("7. Fine-tune Qwen bằng SFT và QLoRA", level=1)
    doc.add_heading("7.1. Cơ chế QLoRA", level=2)
    add_code(doc, "Qwen base nén 4-bit (đóng băng)\n          +\nLoRA adapter nhỏ (được train)\n          ↓\nBase model + adapter → model theo miền ứng dụng")
    add_bullet(doc, "Model gốc giữ nguyên; chỉ các ma trận low-rank được cập nhật.")
    add_bullet(doc, "Adapter nhỏ, dễ version và rollback.")
    add_bullet(doc, "Có thể merge adapter vào model hoặc nạp riêng khi inference.")

    doc.add_heading("7.2. Định dạng SFT hội thoại", level=2)
    add_code(doc, '{\n  "messages": [\n    {"role": "system", "content": "Translate faithfully..."},\n    {"role": "user", "content": "Con không muốn đi học."},\n    {"role": "assistant", "content": "I don\'t want to go to school."}\n  ]\n}')
    add_callout(doc, "Thiết lập quan trọng", "Khi dùng TRL, ưu tiên assistant_only_loss=True để loss chỉ tính trên câu trả lời tiếng Anh, không ép model học lại system prompt và câu người dùng.")

    doc.add_heading("7.3. Điểm khởi đầu cho pilot", level=2)
    add_table(
        doc,
        ["Tham số", "Giá trị khởi đầu", "Ghi chú"],
        [
            ["Base model", "Qwen3-4B-Instruct-2507", "Non-thinking, ổn định cho translation"],
            ["Phương pháp", "QLoRA 4-bit", "Giảm VRAM"],
            ["Max sequence", "512 hoặc 1.024", "Không dùng context 32K/262K"],
            ["Epoch", "2-3", "Dừng sớm nếu validation xấu đi"],
            ["LoRA rank", "16-32", "Bắt đầu nhỏ rồi benchmark"],
            ["GPU", "24 GB", "Cấu hình pilot an toàn"],
        ],
        [2400, 2900, 4060],
    )

    doc.add_heading("7.4. Trình tự một lần train", level=2)
    for index, text in enumerate([
        "Khóa phiên bản base model, tokenizer, dataset và training config.",
        "Kiểm tra dataset schema, trùng lặp, PII và split.",
        "Chạy baseline trên test set trước khi train.",
        "Train adapter và theo dõi training/validation loss.",
        "Chọn checkpoint dựa trên validation, không dựa trên cảm giác.",
        "Chạy test set kín và human review.",
        "Đóng gói adapter, model card nội bộ và báo cáo đánh giá.",
    ], start=1):
        add_numbered(doc, text, index)

    # Evaluation and integration.
    doc.add_heading("8. Đánh giá và tích hợp vào Next.js", level=1)
    doc.add_heading("8.1. Chỉ số chất lượng cần theo dõi", level=2)
    for text in [
        "Giữ đúng phủ định và modality.",
        "Giữ đúng người nói, người nghe và quan hệ gia đình.",
        "Không làm mất màu sắc, số lượng, thời gian và địa điểm.",
        "Không đổi câu kể thành câu hỏi hoặc ngược lại.",
        "Không tự thêm lịch sự hoặc hoàn chỉnh câu thiếu.",
        "Chỉ trả về đúng một câu tiếng Anh, không nhãn hoặc giải thích.",
        "Latency P50/P95, tỷ lệ timeout và tỷ lệ fallback.",
    ]:
        add_check(doc, text)

    doc.add_heading("8.2. Ma trận so sánh bắt buộc", level=2)
    add_table(
        doc,
        ["Ứng viên", "Vai trò", "Điều kiện vượt qua"],
        [
            ["Cloudflare", "Baseline hiện tại", "Đo cùng prompt và test set"],
            ["Qwen nguyên bản", "Xác định có cần train", "Không kém đáng kể ở nhóm lỗi trọng yếu"],
            ["Qwen + LoRA", "Ứng viên production", "Tốt hơn model gốc mà không overfit"],
        ],
        [2400, 3300, 3660],
    )

    doc.add_heading("8.3. Các thay đổi chính trong codebase", level=2)
    add_bullet(doc, "Thêm adapter src/lib/ai/selfHostedText.ts gọi OpenAI-compatible API.")
    add_bullet(doc, "Bổ sung provider self_hosted vào type và telemetry.")
    add_bullet(doc, "Thêm TEXT_AI_PROVIDER để bật/tắt mà không deploy lại model.")
    add_bullet(doc, "Giữ nguyên normalizeTranslationOutput và kiểm tra fidelity hiện có.")
    add_bullet(doc, "Cập nhật test Cloudflare-only thành test provider routing và fallback.")

    doc.add_heading("8.4. Rollout an toàn", level=2)
    add_code(doc, "Offline benchmark → Shadow traffic → Canary 5% → 25% → 50% → 100%")
    add_bullet(doc, "Shadow: Qwen chạy song song nhưng Cloudflare vẫn trả kết quả cho người dùng.")
    add_bullet(doc, "Canary: chỉ một tỷ lệ nhỏ request dùng Qwen; giám sát theo model version.")
    add_bullet(doc, "Circuit breaker: chuyển về Cloudflare khi queue, timeout, OOM hoặc chất lượng bất thường.")

    # Infrastructure and governance.
    doc.add_heading("9. Hạ tầng, bảo mật và vận hành", level=1)
    doc.add_heading("9.1. Cấu hình phần cứng tham khảo", level=2)
    add_table(
        doc,
        ["Mục đích", "GPU/VRAM", "Ghi chú"],
        [
            ["Thử inference INT4", "6-8 GB", "Concurrency thấp"],
            ["Self-host LLM riêng", "12-16 GB", "Đủ cho Qwen 4B và context ngắn"],
            ["Pilot QLoRA", "24 GB", "Cấu hình khuyến nghị"],
            ["LLM + ASR + TTS", "24 GB trở lên", "Phải benchmark concurrency và VRAM"],
        ],
        [2700, 2400, 4260],
    )

    doc.add_heading("9.2. Yêu cầu production", level=2)
    for text in [
        "Docker image và dependency được khóa phiên bản.",
        "HTTPS, API key, firewall và không công khai endpoint GPU.",
        "Health check, readiness check và warm-up model.",
        "Queue, concurrency limit, timeout và request size limit.",
        "Log request ID, model version, adapter version, latency và fallback reason.",
        "Backup adapter, dataset manifest và training config.",
        "Dashboard GPU utilization, VRAM, OOM, queue depth và error rate.",
    ]:
        add_check(doc, text)

    doc.add_heading("9.3. Dữ liệu trẻ em", level=2)
    add_callout(doc, "Yêu cầu bắt buộc", "Phải có consent phù hợp, mục đích sử dụng rõ ràng, thời hạn lưu trữ, quyền xóa dữ liệu và quy trình xử lý sự cố trước khi dùng audio hoặc transcript trẻ em để train.", fill="FFF6E5", accent=GOLD)
    add_bullet(doc, "Dùng speaker_id dạng mã hóa/ẩn danh, không dùng tên thật.")
    add_bullet(doc, "Loại địa chỉ, số điện thoại, trường học và thông tin nhận dạng.")
    add_bullet(doc, "Không đưa dataset riêng tư lên Hugging Face Hub công khai.")
    add_bullet(doc, "Ghi lại consent version và nguồn dữ liệu cho từng bản ghi.")
    add_bullet(doc, "Đặt quyền truy cập tối thiểu cho người gán nhãn và người vận hành.")

    doc.add_heading("9.4. Những rủi ro thường gặp", level=2)
    add_table(
        doc,
        ["Rủi ro", "Cách kiểm soát"],
        [
            ["Dataset sai hoặc nhiễu", "Human review, guideline và audit mẫu"],
            ["Leak train/test", "Dedup theo cụm và split theo người nói/ngữ nghĩa"],
            ["Overfitting", "Validation, early stopping và test kín"],
            ["Model thêm nghĩa", "Fidelity validator và critical-token tests"],
            ["GPU quá tải", "Queue, giới hạn concurrency và fallback"],
            ["Không thể rollback", "Version model/adapter và feature flag"],
        ],
        [3500, 5860],
    )

    # Final checklist and sources.
    doc.add_heading("10. Checklist hành động", level=1)
    doc.add_heading("Giai đoạn A - Chưa train", level=2)
    for text in [
        "Hoàn thành kiến thức Python, token, inference và REST API cơ bản.",
        "Chạy Qwen3-4B-Instruct-2507 bằng vLLM.",
        "Tạo 500-1.000 mẫu golden test có human review.",
        "Benchmark Qwen nguyên bản với Cloudflare.",
        "Quyết định có thật sự cần fine-tune hay không.",
    ]:
        add_check(doc, text)

    doc.add_heading("Giai đoạn B - Fine-tune", level=2)
    for text in [
        "Chuẩn hóa 2.000-5.000 mẫu hội thoại SFT.",
        "Khóa train/validation/test và version dataset.",
        "Chạy QLoRA trên GPU 24 GB.",
        "Đánh giá adapter trên test set kín.",
        "Lưu model card nội bộ, cấu hình và báo cáo kết quả.",
    ]:
        add_check(doc, text)

    doc.add_heading("Giai đoạn C - Tích hợp production", level=2)
    for text in [
        "Thêm provider adapter và biến môi trường.",
        "Cấu hình authentication, timeout, queue và health check.",
        "Chạy shadow traffic trước khi canary.",
        "Giữ Cloudflare fallback và kiểm tra rollback.",
        "Chỉ mở rộng sang ASR/TTS sau khi LLM ổn định.",
    ]:
        add_check(doc, text)

    add_callout(doc, "Mốc tiếp theo nên làm", "Dựng Qwen self-host và benchmark trên 500 câu chuẩn. Đây là bước tạo nhiều thông tin nhất với rủi ro và chi phí thấp nhất.")

    doc.add_heading("11. Tài liệu học chính thức", level=1)
    resources = [
        ("Hugging Face LLM Course", "https://huggingface.co/learn/llm-course/en/chapter1/1", "Transformer, tokenizer, dataset và fine-tuning"),
        ("Qwen3-4B-Instruct-2507", "https://huggingface.co/Qwen/Qwen3-4B-Instruct-2507", "Model card và hướng dẫn sử dụng"),
        ("Qwen3.5-4B", "https://huggingface.co/Qwen/Qwen3.5-4B", "Ứng viên mới hơn để benchmark"),
        ("Qwen + vLLM", "https://github.com/QwenLM/Qwen3/blob/main/docs/source/deployment/vllm.md", "Triển khai OpenAI-compatible API"),
        ("Qwen + Axolotl", "https://github.com/QwenLM/Qwen3/blob/main/docs/source/training/axolotl.md", "SFT, LoRA và QLoRA"),
        ("Hugging Face TRL SFTTrainer", "https://huggingface.co/docs/trl/sft_trainer", "Huấn luyện dataset hội thoại"),
        ("Hugging Face PEFT LoRA", "https://huggingface.co/docs/peft/main/conceptual_guides/lora", "Cơ chế và cấu hình LoRA"),
        ("vLLM OpenAI-compatible server", "https://docs.vllm.ai/en/latest/serving/online_serving/openai_compatible_server/", "API và vận hành model server"),
    ]
    for name, url, description in resources:
        p = doc.add_paragraph()
        set_manual_list_layout(p)
        marker = p.add_run("•\t")
        set_run_font(marker, bold=True, color=BLUE)
        add_hyperlink(p, name, url)
        r = p.add_run(f" - {description}")
        set_run_font(r)

    source_note = doc.add_paragraph(style="Caption")
    source_note.add_run("Ghi chú: Các con số VRAM và thời gian trong tài liệu là ước lượng thực hành, cần được benchmark trên GPU, dataset và mức concurrency thực tế của dự án.")

    # Core properties and final keep rules.
    doc.core_properties.title = "Lộ trình tự học và triển khai AI ngôn ngữ Qwen"
    doc.core_properties.subject = "Cẩm nang cho người mới: dữ liệu, QLoRA, vLLM và tích hợp Next.js"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Qwen, LLM, QLoRA, vLLM, Next.js, Việt-Anh"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())

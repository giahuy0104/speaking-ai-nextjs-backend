from __future__ import annotations

import copy
import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\DELL\Downloads\bao-cao-cloudflare-vs-vertex-ai-bang-chi-phi-moi.docx")
OUTPUT = ROOT / "outputs" / "bao-cao-cloudflare-vs-vertex-ai-quy-mo-5000-user.docx"
BASE_PATH = ROOT / ".codex-tmp" / "create_cloudflare_vs_selfhost_report.py"

spec = importlib.util.spec_from_file_location("base_report", BASE_PATH)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)


USD_VND = 26_000
DAYS = 30
MINUTES_PER_USER_DAY = 10
MINUTES_PER_USER_MONTH = DAYS * MINUTES_PER_USER_DAY
HOURS_PER_USER_MONTH = MINUTES_PER_USER_MONTH / 60
CF_NEURONS_PER_MINUTE = 46.63
CF_FREE_NEURONS_DAY = 10_000
CF_PRICE_PER_1000_NEURONS = 0.011
CF_PAID_BASE = 5.0
CF_NEURONS_USER_DAY = MINUTES_PER_USER_DAY * CF_NEURONS_PER_MINUTE
VERTEX_WARM = 453.33
VERTEX_HA = 906.66
NEW_SCALES = [2000, 3000, 5000]


def money(value: float) -> str:
    return f"${value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def integer_vi(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def vnd_short(value_usd: float) -> str:
    value = value_usd * USD_VND
    if value >= 1_000_000:
        return f"{value / 1_000_000:.2f} triệu".replace(".", ",")
    if value >= 1_000:
        return f"{value / 1_000:.0f} nghìn"
    return f"{value:.0f}đ"


def cloudflare_total(users: int) -> float:
    excess_day = max(0.0, CF_NEURONS_USER_DAY * users - CF_FREE_NEURONS_DAY)
    ai_cost = excess_day * DAYS * CF_PRICE_PER_1000_NEURONS / 1000
    return 0.0 if excess_day == 0 else CF_PAID_BASE + ai_cost


def replace_tc_text(tr, cell_index: int, text: str) -> None:
    cells = tr.findall(qn("w:tc"))
    cell = cells[cell_index]
    text_nodes = cell.findall(".//" + qn("w:t"))
    if text_nodes:
        text_nodes[0].text = text
        for extra in text_nodes[1:]:
            extra.text = ""
        return

    paragraph = cell.find(qn("w:p"))
    if paragraph is None:
        paragraph = OxmlElement("w:p")
        cell.append(paragraph)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    paragraph.append(run)


def clone_data_row(template_row, values: list[str]):
    tr = copy.deepcopy(template_row._tr)
    for idx, value in enumerate(values):
        if value is not None:
            replace_tc_text(tr, idx, value)
    return tr


def compact_table(table, body_size: float, vertical_margin: int, preserve_group_column: bool = False) -> None:
    for ridx, row in enumerate(table.rows):
        base.prevent_row_split(row)
        for cidx, cell in enumerate(row.cells):
            base.set_cell_margins(cell, top=vertical_margin, bottom=vertical_margin, start=120, end=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            if ridx > 0 and not (preserve_group_column and cidx == 0):
                base.set_cell_shading(cell, "F8FAFC" if ridx % 2 == 0 else "FFFFFF")

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = base.Pt(0)
                paragraph.paragraph_format.space_after = base.Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if preserve_group_column and ridx > 0 and cidx == 0:
                    continue
                for run in paragraph.runs:
                    base.set_run_font(
                        run,
                        size=body_size + 0.3 if ridx == 0 else body_size,
                        color=base.NAVY if ridx == 0 else base.INK,
                        bold=True if ridx == 0 else None,
                    )


def add_cloudflare_rows(table, template_row, before_row=None) -> None:
    for users in NEW_SCALES:
        total = cloudflare_total(users)
        values = [
            None,
            integer_vi(users),
            integer_vi(MINUTES_PER_USER_MONTH * users),
            money(total),
            vnd_short(total),
        ]
        tr = clone_data_row(template_row, values)
        if before_row is None:
            table._tbl.append(tr)
        else:
            before_row._tr.addprevious(tr)


def edit_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    if len(doc.tables) != 3:
        raise RuntimeError(f"Expected 3 tables, found {len(doc.tables)}")

    # Table 1: grouped Cloudflare and Vertex cost summary.
    summary = doc.tables[0]
    vertex_start = summary.rows[5]
    cloudflare_template = summary.rows[4]
    for users in NEW_SCALES:
        total = cloudflare_total(users)
        tr = clone_data_row(
            cloudflare_template,
            [None, integer_vi(users), integer_vi(MINUTES_PER_USER_MONTH * users), money(total), vnd_short(total)],
        )
        vertex_start._tr.addprevious(tr)

    vertex_template = summary.rows[-1]
    for users in NEW_SCALES:
        tr = clone_data_row(
            vertex_template,
            [None, integer_vi(users), integer_vi(MINUTES_PER_USER_MONTH * users), money(VERTEX_WARM), vnd_short(VERTEX_WARM)],
        )
        summary._tbl.append(tr)
    compact_table(summary, body_size=7.8, vertical_margin=36, preserve_group_column=True)

    # Table 2: Cloudflare detailed costs.
    cloudflare = doc.tables[1]
    cloudflare_template = cloudflare.rows[-1]
    for users in NEW_SCALES:
        total = cloudflare_total(users)
        tr = clone_data_row(
            cloudflare_template,
            [integer_vi(users), integer_vi(MINUTES_PER_USER_MONTH * users), money(total), vnd_short(total)],
        )
        cloudflare._tbl.append(tr)
    compact_table(cloudflare, body_size=8.6, vertical_margin=52)

    # Table 3: Vertex always-on replica costs.
    vertex = doc.tables[2]
    vertex_template = vertex.rows[-1]
    for users in NEW_SCALES:
        tr = clone_data_row(
            vertex_template,
            [integer_vi(users), integer_vi(int(HOURS_PER_USER_MONTH * users)), money(VERTEX_WARM), money(VERTEX_HA)],
        )
        vertex._tbl.append(tr)
    compact_table(vertex, body_size=8.3, vertical_margin=48)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    edit_document()

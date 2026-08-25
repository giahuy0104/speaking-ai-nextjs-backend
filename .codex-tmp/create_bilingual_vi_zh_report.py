from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\DELL\Downloads\bao-cao-cloudflare-vs-vertex-ai-quy-mo-5000-user.docx")
OUTPUT = ROOT / "outputs" / "bao-cao-cloudflare-vs-vertex-ai-quy-mo-5000-user-viet-trung.docx"

ARIAL = "Arial"
CHINESE_FONT = "Microsoft YaHei"
NAVY = "17365D"
BLUE = "2E74B5"
INK = "1F2937"
MUTED = "667085"
GREEN = "276749"
GOLD = "7A5A00"


TRANSLATIONS = {
    0: "部署 AI 模型：选择 Cloudflare 还是 Vertex AI？",
    1: "按每位用户每天说话 10 分钟估算的成本",
    2: "每位用户的用量：每天 10 分钟 × 30 天 = 每月 300 分钟 = 5 小时音频",
    3: "Cloudflare：Workers AI + @cf/openai/whisper-large-v3-turbo",
    4: "Vertex：自定义端点：n1-standard-4 + NVIDIA Tesla T4（us-central1）",
    5: "价格日期：2026 年 8 月 21 日；按 1 美元 = 26,000 越南盾换算；未含税",
    6: "结论：",
    7: "在 1 至 1,000 名用户的范围内，默认推荐 Cloudflare：成本更低、部署更快，并且无需运维 GPU。",
    8: "只有在自定义模型能够显著提升质量、必须使用 GCP 的 IAM/VPC，或实际基准测试证明总成本更低时，才切换到 Vertex AI。",
    9: "不同用户规模的成本",
    11: "1. Cloudflare Workers AI",
    12: "计价说明：",
    13: "每位用户每天使用 466.3 Neurons。",
    14: "每月 AI 成本 = max(0；466.3 × 用户数 - 10,000) × 30 × 0.011 / 1,000。",
    15: "当每日用量超过免费额度时，需要至少支付每月 5 美元的 Workers Paid 套餐费。",
    16: "使用 Workers Paid 后，超出免费额度的 AI 用量按每 1,000 Neurons 0.011 美元另行计费。",
    18: "2. Vertex AI 自定义端点",
    19: "参考单价：",
    20: "1 个副本（服务器）的总成本：每小时 0.621 美元。",
    21: "1 个副本（服务器）全天候运行 = 每月 453.33 美元。",
    22: "音频小时/月：一个月内累计的输入音频总时长，并不表示 GPU 一定要运行整整 5 小时。",
    23: "1 个副本（1 台服务器）：服务器始终开启，Whisper 模型已预先加载到 GPU。",
    24: "2 个高可用副本（2 台服务器）：系统将同时使用 2 台服务器运行：",
    25: "如果一台服务器发生故障，另一台仍可继续提供服务。",
    26: "当多位用户同时发送音频时分担负载。",
}


PARAGRAPH_STYLE = {
    0: (16.0, NAVY, True, False),
    1: (10.8, MUTED, False, False),
    2: (8.8, MUTED, False, False),
    3: (8.8, MUTED, False, False),
    4: (8.8, MUTED, False, False),
    5: (8.8, MUTED, False, False),
    6: (9.2, GREEN, True, False),
    7: (8.8, INK, False, False),
    8: (8.8, INK, False, False),
    9: (11.0, BLUE, True, False),
    11: (11.0, BLUE, True, False),
    12: (9.2, BLUE, True, False),
    13: (8.7, MUTED, False, False),
    14: (8.7, MUTED, False, False),
    15: (8.7, MUTED, False, False),
    16: (8.7, MUTED, False, False),
    18: (11.0, BLUE, True, False),
    19: (9.2, BLUE, True, False),
    20: (8.8, MUTED, False, False),
    21: (8.8, MUTED, False, False),
    22: (8.2, MUTED, False, True),
    23: (8.2, MUTED, False, True),
    24: (8.2, MUTED, False, True),
    25: (8.2, INK, False, False),
    26: (8.2, INK, False, False),
}


def set_run_font(run, size: float, color: str, bold: bool = False, italic: bool = False, chinese: bool = False) -> None:
    run.font.name = CHINESE_FONT if chinese else ARIAL
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ARIAL)
    rfonts.set(qn("w:hAnsi"), ARIAL)
    rfonts.set(qn("w:eastAsia"), CHINESE_FONT)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def append_translation(paragraph, text: str, size: float, color: str, bold: bool, italic: bool) -> None:
    break_run = paragraph.add_run()
    break_run.add_break()
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic, chinese=True)
    paragraph.paragraph_format.line_spacing = 1.0


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def set_bilingual_cell(
    cell,
    vietnamese: str,
    chinese: str,
    vi_size: float,
    zh_size: float,
    color: str,
    bold: bool = True,
    alignment=WD_ALIGN_PARAGRAPH.CENTER,
) -> None:
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0

    vi_run = paragraph.add_run(vietnamese)
    set_run_font(vi_run, size=vi_size, color=color, bold=bold)
    vi_run.add_break()
    zh_run = paragraph.add_run(chinese)
    set_run_font(zh_run, size=zh_size, color=color, bold=bold, chinese=True)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top: int, bottom: int, start: int = 120, end: int = 120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def translate_tables(doc: Document) -> None:
    summary, cloudflare, vertex = doc.tables

    summary_headers = [
        ("Giải pháp", "方案"),
        ("User", "用户数"),
        ("Phút/tháng", "分钟/月"),
        ("Tổng USD", "合计（美元）"),
        ("Quy đổi VND", "折合越南盾"),
    ]
    for idx, (vi, zh) in enumerate(summary_headers):
        set_bilingual_cell(
            summary.rows[0].cells[idx], vi, zh, 7.8, 7.0, NAVY,
            alignment=WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_margins(summary.rows[0].cells[idx], 55, 55)

    cloudflare_group = summary.rows[1].cells[0]
    vertex_group = summary.rows[8].cells[0]
    set_bilingual_cell(cloudflare_group, "Cloudflare Workers AI", "Cloudflare Workers AI 方案", 7.8, 7.0, GREEN, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    set_bilingual_cell(vertex_group, "Vertex AI", "Vertex AI 方案", 7.8, 7.0, GOLD, alignment=WD_ALIGN_PARAGRAPH.LEFT)

    cloudflare_headers = [
        ("User", "用户数"),
        ("Phút/tháng", "分钟/月"),
        ("Tổng USD", "合计（美元）"),
        ("Quy đổi VND", "折合越南盾"),
    ]
    for idx, (vi, zh) in enumerate(cloudflare_headers):
        set_bilingual_cell(cloudflare.rows[0].cells[idx], vi, zh, 8.1, 7.3, NAVY)
        set_cell_margins(cloudflare.rows[0].cells[idx], 65, 65)

    vertex_headers = [
        ("User", "用户数"),
        ("Audio giờ/tháng", "音频小时/月"),
        ("1 replica(1 máy chủ)", "1 个副本（1 台服务器）"),
        ("2 replica HA(2 máy chủ)", "2 个高可用副本（2 台服务器）"),
    ]
    for idx, (vi, zh) in enumerate(vertex_headers):
        set_bilingual_cell(vertex.rows[0].cells[idx], vi, zh, 7.6, 6.9, NAVY)
        set_cell_margins(vertex.rows[0].cells[idx], 55, 55)


def translate_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header.paragraphs[0]
        if header.runs:
            header.runs[0].text = "BÁO CÁO QUYẾT ĐỊNH CÔNG NGHỆ / 技术决策报告"
            set_run_font(header.runs[0], size=7.7, color=MUTED, bold=True, chinese=True)
        if len(header.runs) > 1:
            set_run_font(header.runs[1], size=8.0, color=MUTED)

        footer = section.footer.paragraphs[0]
        if footer.runs:
            footer.runs[0].text = "Trang / 页 "
            set_run_font(footer.runs[0], size=8.0, color=MUTED, chinese=True)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    if len(doc.paragraphs) < 27 or len(doc.tables) != 3:
        raise RuntimeError("The source document structure does not match the expected report.")

    for index, chinese in TRANSLATIONS.items():
        size, color, bold, italic = PARAGRAPH_STYLE[index]
        append_translation(doc.paragraphs[index], chinese, size, color, bold, italic)

    # Keep the executive summary on page 1 and place the full scale-cost table
    # on a clean page instead of leaving a short table continuation page.
    doc.paragraphs[9].paragraph_format.page_break_before = True

    translate_tables(doc)
    translate_header_footer(doc)

    props = doc.core_properties
    props.title = "Báo cáo Cloudflare và Vertex AI song ngữ Việt - Trung"
    props.subject = "Chi phí ASR theo quy mô người dùng / 按用户规模估算 ASR 成本"
    props.keywords = "Cloudflare, Vertex AI, ASR, Việt Trung, 越中双语"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

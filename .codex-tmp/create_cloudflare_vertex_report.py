from __future__ import annotations

import importlib.util
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
BASE_PATH = ROOT / ".codex-tmp" / "create_cloudflare_vs_selfhost_report.py"
OUTPUT_DIR = ROOT / "outputs"
TMP_DIR = ROOT / ".codex-tmp" / "cloudflare-vertex-report"
OUTPUT_PATH = OUTPUT_DIR / "bao-cao-so-sanh-cloudflare-vertex-ai-va-ai-tu-van-hanh.docx"
CHART_PATH = TMP_DIR / "chi-phi-asr-tren-moi-nguoi-dung.png"

spec = importlib.util.spec_from_file_location("base_report", BASE_PATH)
base = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(base)


USD_VND = 26_000
VERTEX_MACHINE_PER_HOUR = 0.219
VERTEX_T4_PER_HOUR = 0.402
VERTEX_NODE_PER_HOUR = VERTEX_MACHINE_PER_HOUR + VERTEX_T4_PER_HOUR
HOURS_PER_MONTH = 730
VERTEX_WARM_MONTH = VERTEX_NODE_PER_HOUR * HOURS_PER_MONTH


def money(value: float, digits: int = 2) -> str:
    return f"${value:,.{digits}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def create_chart(path: Path) -> None:
    width, height = 1400, 720
    img = Image.new("RGB", (width, height), f"#{base.WHITE}")
    draw = ImageDraw.Draw(img)
    font_regular = "C:/Windows/Fonts/arial.ttf"
    font_bold = "C:/Windows/Fonts/arialbd.ttf"
    title_font = ImageFont.truetype(font_bold, 40)
    label_font = ImageFont.truetype(font_regular, 25)
    small_font = ImageFont.truetype(font_regular, 22)
    bold_font = ImageFont.truetype(font_bold, 23)

    draw.text((70, 32), "Chi phí ASR trên mỗi người dùng mỗi tháng", font=title_font, fill=f"#{base.NAVY}")
    draw.text(
        (70, 87),
        "10 phút/ngày; Vertex: 1 endpoint T4 luôn sẵn sàng; chưa gồm thuế và nhân sự",
        font=small_font,
        fill=f"#{base.MUTED}",
    )

    users = [100, 500, 1000]
    cloudflare_totals = [18.65, 79.85, 156.35]
    cloudflare = [total / user for total, user in zip(cloudflare_totals, users)]
    vertex = [VERTEX_WARM_MONTH / u for u in users]
    speech = [4.80, 4.80, 4.80]
    series = [
        ("Cloudflare", cloudflare, base.BLUE),
        ("Vertex T4 warm", vertex, base.NAVY),
        ("Google STT realtime", speech, base.GOLD),
    ]

    max_y = 5.2
    left, top, right, bottom = 115, 145, 1335, 585
    chart_h = bottom - top
    chart_w = right - left
    for value in range(0, 6):
        y = bottom - (value / max_y) * chart_h
        draw.line((left, y, right, y), fill="#D8DEE8", width=2)
        label = f"${value}"
        bbox = draw.textbbox((0, 0), label, font=small_font)
        draw.text((left - 16 - (bbox[2] - bbox[0]), y - 12), label, font=small_font, fill=f"#{base.MUTED}")

    group_width = chart_w / 3
    bar_width = 78
    offsets = [-88, 0, 88]
    for idx, user_count in enumerate(users):
        center = left + group_width * (idx + 0.5)
        for sidx, (_, values, color) in enumerate(series):
            value = values[idx]
            x0 = center + offsets[sidx] - bar_width / 2
            x1 = center + offsets[sidx] + bar_width / 2
            y0 = bottom - (value / max_y) * chart_h
            draw.rounded_rectangle((x0, y0, x1, bottom), radius=7, fill=f"#{color}")
            label = f"${value:.2f}"
            bbox = draw.textbbox((0, 0), label, font=bold_font)
            draw.text(((x0 + x1 - (bbox[2] - bbox[0])) / 2, y0 - 30), label, font=bold_font, fill=f"#{color}")
        user_text = f"{user_count:,}".replace(",", ".") + " người"
        bbox = draw.textbbox((0, 0), user_text, font=label_font)
        draw.text((center - (bbox[2] - bbox[0]) / 2, bottom + 15), user_text, font=label_font, fill=f"#{base.INK}")

    legend_y = 666
    starts = [240, 545, 920]
    for start, (label, _, color) in zip(starts, series):
        draw.rounded_rectangle((start, legend_y - 9, start + 28, legend_y + 19), radius=4, fill=f"#{color}")
        draw.text((start + 40, legend_y - 13), label, font=small_font, fill=f"#{base.INK}")
    img.save(path, quality=95)


def add_caption_with_links(doc, items: list[tuple[str, str]]) -> None:
    p = doc.add_paragraph(style="Caption")
    p.add_run("Nguồn: ")
    for idx, (label, url) in enumerate(items):
        if idx:
            p.add_run("; ")
        base.add_hyperlink(p, label, url)
    p.add_run(".")


def build_report() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    create_chart(CHART_PATH)

    doc = Document()
    base.configure_styles(doc)
    base.configure_sections(doc)
    base.BULLET_NUM_ID, base.NUMBER_NUM_ID = base.add_custom_numbering(doc)

    props = doc.core_properties
    props.title = "So sánh Cloudflare, Vertex AI và model AI tự vận hành"
    props.subject = "Độ khó, chi phí và ngưỡng chuyển đổi với 10 phút audio/người/ngày"
    props.author = "Speaking AI"
    props.last_modified_by = "Speaking AI"
    props.keywords = "Cloudflare Workers AI, Vertex AI, Speech-to-Text, Whisper, self-host, ASR, chi phí"
    props.created = datetime(2026, 8, 21, tzinfo=timezone.utc)

    # Trang 1: Tóm tắt điều hành.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(0.18)
    p.paragraph_format.space_after = Inches(0.04)
    base.set_run_font(p.add_run("ĐÁNH GIÁ CHIẾN LƯỢC ASR"), size=10.5, color=base.BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Inches(0.06)
    base.set_run_font(p.add_run("Cloudflare, Vertex AI\nhay model tự vận hành?"), size=25, color=base.NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Inches(0.15)
    base.set_run_font(p.add_run("So sánh độ khó, chi phí và thời điểm dừng từng giải pháp"), size=14, color=base.MUTED)

    metadata = [
        ("Phạm vi", "100, 500 và 1.000 người dùng; mỗi người nói 10 phút/ngày"),
        ("Bài toán", "Nhận dạng giọng nói (ASR) thời gian thực cho ứng dụng hội thoại"),
        ("Ngày giá", "21/08/2026; quy đổi 26.000 VND/USD; chưa gồm thuế"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Inches(0.025)
        base.set_run_font(p.add_run(f"{label}: "), size=10.5, color=base.NAVY, bold=True)
        base.set_run_font(p.add_run(value), size=10.5, color=base.INK)

    base.add_callout(
        doc,
        "KHUYẾN NGHỊ",
        "Giữ Cloudflare làm production primary ở quy mô 100-1.000 người dùng. Chỉ POC Vertex AI custom endpoint khi cần model tùy biến, IAM/VPC của Google Cloud hoặc chất lượng giọng trẻ em tốt hơn. Không nên chọn Google Speech-to-Text hay endpoint Vertex luôn chạy chỉ để giảm chi phí tại quy mô này.",
        fill=base.LIGHT_BLUE,
        accent=base.NAVY,
    )
    base.add_bullet(doc, "Cloudflare: khoảng 17,00-20,30 USD/tháng ở 100 user và 154,70-158,00 USD ở 1.000 user.", base.BULLET_NUM_ID, "Cloudflare:")
    base.add_bullet(doc, "Vertex AI T4 luôn sẵn sàng: khoảng 453,33 USD/tháng cho một replica; khoảng 906,66 USD cho hai replica.", base.BULLET_NUM_ID, "Vertex AI T4 luôn sẵn sàng:")
    base.add_bullet(doc, "Google Speech-to-Text realtime: 480-4.800 USD/tháng; đây là API Google được quản lý hoàn toàn, không phải custom endpoint.", base.BULLET_NUM_ID, "Google Speech-to-Text realtime:")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Inches(0.02)
    inline = p.add_run().add_picture(str(CHART_PATH), width=Inches(5.65))
    base.set_picture_alt_text(
        inline,
        "Chi phí ASR trên mỗi người dùng",
        "Biểu đồ so sánh Cloudflare, một Vertex AI T4 luôn sẵn sàng và Google Speech-to-Text realtime tại 100, 500 và 1.000 người dùng.",
    )
    cap = doc.add_paragraph("Hình 1. Cloudflare có chi phí trên mỗi user thấp nhất ở cả ba mốc.", style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Trang 2: Phạm vi và định nghĩa.
    base.add_page_break(doc)
    base.add_heading(doc, "1. Phạm vi, giả định và cách hiểu Vertex AI", 1)
    base.add_body(
        doc,
        "Bảng chính chỉ tính ASR. Số phút nói không đủ để suy ra chính xác token của model hội thoại hoặc số ký tự TTS. Chi phí lưu trữ, log, egress, thuế và nhân sự được nêu là khoản cộng thêm, không gộp vào con số ASR.",
    )
    base.add_table(
        doc,
        ["Người dùng", "Audio/ngày", "Audio/tháng", "Giờ audio/tháng"],
        [
            ["100", "1.000 phút", "30.000 phút", "500 giờ"],
            ["500", "5.000 phút", "150.000 phút", "2.500 giờ"],
            ["1.000", "10.000 phút", "300.000 phút", "5.000 giờ"],
        ],
        [1500, 2400, 2760, 2700],
        [WD_ALIGN_PARAGRAPH.CENTER] * 4,
        font_size=9.6,
    )

    base.add_heading(doc, "Bốn phương án được phân biệt rõ", 2)
    base.add_table(
        doc,
        ["Phương án", "Bản chất", "Cách tính phí", "Phù hợp nhất khi"],
        [
            ["Cloudflare Workers AI", "Whisper được quản lý", "Theo phút audio + Workers", "Cần ra nhanh, chi phí thấp"],
            ["Vertex AI custom endpoint", "Tự đóng gói Whisper/model riêng", "Máy + GPU theo giờ", "Cần model tùy biến và MLOps GCP"],
            ["Google Speech-to-Text", "ASR Google được quản lý hoàn toàn", "Theo phút audio", "Ưu tiên Google ASR và governance"],
            ["Self-host", "Tự vận hành GPU và model", "GPU, máy, điện, nhân sự", "Cần toàn quyền kiểm soát"],
        ],
        [2050, 2550, 1980, 2780],
        [WD_ALIGN_PARAGRAPH.LEFT] * 4,
        font_size=8.7,
        header_fill=base.LIGHT_BLUE,
    )
    base.add_callout(
        doc,
        "LƯU Ý",
        "Vertex AI không tự động đồng nghĩa với Speech-to-Text. Đưa Whisper/PhoWhisper lên Vertex là một dịch vụ custom prediction; gọi Speech-to-Text là một API riêng với giá theo phút.",
        fill=base.PALE_GOLD,
        accent=base.GOLD,
    )

    base.add_heading(doc, "2. Chi phí Cloudflare", 1)
    base.add_body(doc, "Cloudflare Whisper Large V3 Turbo có giá 0,00051 USD/phút. Workers Paid tối thiểu 5 USD/tháng; quota Workers AI miễn phí có giá trị tối đa khoảng 3,30 USD/tháng và dùng chung giữa các model.")
    base.add_table(
        doc,
        ["User", "ASR", "Tổng Cloudflare", "Quy đổi VND", "Chi phí/user"],
        [
            ["100", "$15,30", "$17,00-$20,30", "442.000-528.000đ", "$0,17-$0,20"],
            ["500", "$76,50", "$78,20-$81,50", "2,03-2,12 triệu", "$0,16"],
            ["1.000", "$153,00", "$154,70-$158,00", "4,02-4,11 triệu", "$0,15-$0,16"],
        ],
        [1050, 1300, 2100, 2550, 2360],
        [WD_ALIGN_PARAGRAPH.CENTER] * 5,
        font_size=9.0,
        header_fill=base.PALE_GREEN,
    )
    add_caption_with_links(doc, [
        ("Cloudflare Whisper pricing", "https://developers.cloudflare.com/workers-ai/models/whisper-large-v3-turbo/"),
        ("Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/"),
        ("Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/"),
    ])

    # Trang 3: Vertex.
    base.add_page_break(doc)
    base.add_heading(doc, "3. Chi phí triển khai custom model trên Vertex AI", 1)
    base.add_body(
        doc,
        "Kịch bản tham chiếu dùng một n1-standard-4 (0,219 USD/giờ) và một NVIDIA Tesla T4 (0,402 USD/giờ) tại Iowa/us-central1. Tổng là 0,621 USD/giờ hoặc khoảng 453,33 USD/tháng nếu replica chạy 730 giờ. Giá thực tế thay đổi theo vùng, loại máy, storage, log và egress.",
    )
    base.add_callout(
        doc,
        "CÔNG THỨC",
        "Vertex warm = (0,219 + 0,402) x 730 giờ = 453,33 USD/tháng/replica. Hai replica để tăng dự phòng = 906,66 USD/tháng, chưa gồm chi phí phụ.",
        fill=base.LIGHT_GRAY,
        accent=base.BLUE,
    )
    base.add_table(
        doc,
        ["User", "1 replica warm", "2 replica HA", "Quy đổi 1 replica", "So với Cloudflare"],
        [
            ["100", "$453,33", "$906,66", "11,79 triệu", "Khoảng 22-27 lần"],
            ["500", "$453,33", "$906,66", "11,79 triệu", "Khoảng 5,6-5,8 lần"],
            ["1.000", "$453,33", "$906,66", "11,79 triệu", "Khoảng 2,9 lần"],
        ],
        [1050, 1670, 1670, 2300, 2670],
        [WD_ALIGN_PARAGRAPH.CENTER] * 5,
        font_size=9.0,
        header_fill=base.LIGHT_BLUE,
    )

    base.add_heading(doc, "3.1. Kịch bản scale-to-zero", 2)
    base.add_body(
        doc,
        "Vertex cho phép đặt min replica bằng 0 ở tính năng scale-to-zero. Tuy nhiên tài liệu CLI hiện ghi tính năng này ở beta, mặc định scale xuống sau 1 giờ không có traffic và giữ scale-up tối thiểu 1 giờ. Với người dùng nói rải rác cả ngày, endpoint có thể gần như luôn nóng; do đó bảng dưới chỉ là compute thuần lý thuyết.",
    )
    base.add_table(
        doc,
        ["User", "GPU giờ ở tốc độ 10-30x realtime", "Compute Vertex lý thuyết", "Rủi ro"],
        [
            ["100", "16,7-50 giờ", "$10,35-$31,05", "Cold start; request rải rác làm tăng giờ"],
            ["500", "83,3-250 giờ", "$51,75-$155,25", "Batching quyết định hiệu quả"],
            ["1.000", "166,7-500 giờ", "$103,50-$310,50", "Phải benchmark tải đỉnh"],
        ],
        [1050, 2900, 2300, 3110],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.9,
        header_fill=base.PALE_GOLD,
    )
    base.add_bullet(doc, "Ưu điểm Vertex: Model Registry, IAM, endpoint, autoscaling, logging, VPC và pipeline trong cùng hệ sinh thái GCP.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Việc vẫn phải làm: container hóa model, xử lý audio, batching, benchmark GPU, monitoring chất lượng và rollback.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Một T4 không được mặc định là đủ tải cho 1.000 user; số phiên đồng thời và p95 latency phải được đo bằng audio thật.", base.BULLET_NUM_ID)
    add_caption_with_links(doc, [
        ("Vertex/Agent Platform pricing", "https://cloud.google.com/products/gemini-enterprise-agent-platform/pricing"),
        ("Scale-to-zero deployment flags", "https://docs.cloud.google.com/sdk/gcloud/reference/beta/ai/endpoints/deploy-model"),
        ("Custom container for inference", "https://docs.cloud.google.com/vertex-ai/docs/predictions/use-custom-container"),
    ])

    # Trang 4: Google STT và chi phí tổng hợp.
    base.add_page_break(doc)
    base.add_heading(doc, "4. Google Speech-to-Text: phương án Google dễ triển khai hơn", 1)
    base.add_body(
        doc,
        "Speech-to-Text V2 Standard tính 0,016 USD/phút cho 0-500.000 phút/tháng. Dynamic Batch là 0,003 USD/phút nhưng có độ ưu tiên thấp và có thể chờ tới 24 giờ, nên không phù hợp cho hội thoại thời gian thực.",
    )
    base.add_table(
        doc,
        ["User", "Realtime Standard", "Quy đổi VND", "Dynamic Batch", "Nhận xét"],
        [
            ["100", "$480", "12,48 triệu", "$90", "Dễ dùng nhưng đắt hơn Cloudflare"],
            ["500", "$2.400", "62,40 triệu", "$450", "Dynamic Batch không realtime"],
            ["1.000", "$4.800", "124,80 triệu", "$900", "Cần lý do chất lượng/compliance"],
        ],
        [1050, 1800, 1850, 1750, 2910],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.9,
        header_fill=base.LIGHT_BLUE,
    )
    base.add_body(doc, "Có mức Standard Logged 0,012 USD/phút nếu chấp nhận điều kiện logging dữ liệu; cần đánh giá riêng về dữ liệu trẻ em, consent và chính sách nội bộ trước khi dùng.")
    add_caption_with_links(doc, [("Google Cloud Speech-to-Text pricing", "https://cloud.google.com/speech-to-text/pricing")])

    base.add_heading(doc, "5. Tổng hợp chi phí production tham chiếu", 1)
    base.add_table(
        doc,
        ["Phương án", "100 user", "500 user", "1.000 user", "Độ khó"],
        [
            ["Cloudflare", "$17-$20", "$78-$82", "$155-$158", "Thấp (2/10)"],
            ["Vertex 1 T4 warm", "$453", "$453", "$453", "Trung bình-cao (7/10)"],
            ["Google STT realtime", "$480", "$2.400", "$4.800", "Thấp-trung bình (3/10)"],
            ["GPU cloud self-host 24/7", "$497-$929", "$497-$929", "$497-$929", "Cao (8/10)"],
            ["Máy GPU tự mua", "$150-$300", "$150-$300", "$150-$300", "Rất cao (9/10)"],
        ],
        [2400, 1450, 1600, 1750, 2160],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        font_size=8.6,
        header_fill=base.PALE_GREEN,
    )
    base.add_body(doc, "Các con số self-host là khoảng tham chiếu trước nhân sự. Máy tự mua đã gồm khấu hao, điện/mạng/storage ở mức ước tính nhưng chưa gồm dự phòng và công vận hành.", italic=True, color=base.MUTED)

    # Trang 5: Độ khó.
    base.add_page_break(doc)
    base.add_heading(doc, "6. So sánh độ khó và năng lực cần có", 1)
    base.add_table(
        doc,
        ["Tiêu chí", "Cloudflare", "Vertex custom", "Google STT", "Self-host"],
        [
            ["POC ước tính", "1-3 ngày", "2-4 tuần", "2-5 ngày", "3-8 tuần"],
            ["GPU/CUDA", "Không", "Cần hiểu", "Không", "Bắt buộc"],
            ["Autoscaling", "Được quản lý", "Cấu hình + benchmark", "Được quản lý", "Tự xây"],
            ["Model tùy biến", "Hạn chế", "Cao", "Theo tính năng API", "Rất cao"],
            ["IAM/VPC/audit", "Theo Cloudflare", "Mạnh trong GCP", "Mạnh trong GCP", "Tự thiết kế"],
            ["Cold start", "Thấp", "Có nếu scale-to-zero", "Thấp", "Tự xử lý"],
            ["HA production", "Đơn giản", "Nên >=2 replica", "Đơn giản", "Tự xây >=2 node"],
            ["Nhân sự vận hành", "Thấp", "Trung bình", "Thấp", "Cao"],
        ],
        [2070, 1740, 2060, 1740, 1750],
        [WD_ALIGN_PARAGRAPH.LEFT] * 5,
        font_size=8.3,
        header_fill=base.LIGHT_BLUE,
    )
    base.add_callout(
        doc,
        "SLA",
        "Google Cloud công bố SLO 99,5% cho Custom Model Online Prediction khi model triển khai trên từ hai node. Một replica rẻ hơn nhưng không nên được xem là cấu hình HA production.",
        fill=base.PALE_GOLD,
        accent=base.GOLD,
    )
    add_caption_with_links(doc, [("Vertex AI Service Level Agreement", "https://cloud.google.com/vertex-ai/sla")])

    base.add_heading(doc, "Phần khó nhất khi chọn Vertex hoặc self-host", 2)
    base.add_bullet(doc, "Xây golden set đại diện cho giọng trẻ em, vùng miền, thiết bị và tiếng ồn thực tế.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Đo WER, lỗi ngữ nghĩa, p50/p95 latency và số phiên đồng thời thay vì chỉ đo tốc độ một file dài.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Tối ưu đoạn audio ngắn, batching, tải model, retry và timeout mà không làm hội thoại bị trễ.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Thiết lập canary, rollback, theo dõi drift, quyền dữ liệu và quy trình fine-tune có thể lặp lại.", base.BULLET_NUM_ID)

    # Trang 6: Điểm hòa vốn và stop rules.
    base.add_page_break(doc)
    base.add_heading(doc, "7. Điểm hòa vốn và thời điểm dừng giải pháp", 1)
    base.add_callout(
        doc,
        "ĐIỂM HÒA VỐN",
        "Một user tạo 300 phút/tháng, tương đương khoảng 0,153 USD ASR Cloudflare. Một Vertex T4 warm 453,33 USD/tháng bắt đầu gần hòa vốn hạ tầng với Cloudflare ở khoảng 2.900-3.000 user. Hai replica HA đẩy ngưỡng lên khoảng 5.900 user, trước storage, logging, egress và nhân sự.",
        fill=base.LIGHT_GRAY,
        accent=base.BLUE,
    )

    base.add_heading(doc, "Khi nào dừng Cloudflare làm primary", 2)
    cloudflare_stop = [
        ("Chi phí", "TCO của Vertex/self-host thấp hơn Cloudflare ít nhất 20% trong 3 tháng liên tiếp."),
        ("Chất lượng", "custom model giảm WER hoặc lỗi hiểu sai ít nhất 10-15% trên golden set."),
        ("Độ trễ", "p95 đạt mục tiêu trải nghiệm ở tải đỉnh gấp đôi dự kiến."),
        ("Ổn định", "shadow/canary tối thiểu 30 ngày, có rollback và fallback đã thử."),
        ("Dữ liệu", "quyền sử dụng, consent và quy trình lưu/xóa dữ liệu đã được duyệt."),
    ]
    for label, detail in cloudflare_stop:
        base.add_numbered(doc, f"{label}: {detail}", base.NUMBER_NUM_ID, f"{label}:")

    base.add_heading(doc, "Khi nào dừng POC Vertex AI", 2)
    base.add_bullet(doc, "Không cải thiện chất lượng ít nhất 10% sau hai vòng tuning hoặc dữ liệu không đủ tốt.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "GPU utilization dưới 30% và traffic phân tán khiến scale-to-zero không tạo lợi ích thực tế.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Chi phí đầy đủ không thấp hơn Cloudflare 20%, hoặc cần hai replica nhưng quy mô dưới khoảng 5.900 user.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Cold start/p95 latency không đạt, hoặc nhóm phải dành hơn một ngày kỹ thuật mỗi tuần để vận hành.", base.BULLET_NUM_ID)

    base.add_heading(doc, "Khi nào dừng Google Speech-to-Text", 2)
    base.add_bullet(doc, "Không có lợi thế đo được về chất lượng, compliance hoặc hệ sinh thái GCP để bù mức 0,016 USD/phút.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Chi phí vượt ngân sách; Dynamic Batch không thể thay thế realtime cho luồng hội thoại.", base.BULLET_NUM_ID)

    base.add_heading(doc, "Khi nào dừng self-host", 2)
    base.add_bullet(doc, "Sự cố GPU/CUDA, HA hoặc bảo mật chiếm quá nhiều thời gian của đội; chất lượng không vượt model managed.", base.BULLET_NUM_ID)
    base.add_bullet(doc, "Tổng chi phí máy, dự phòng và nhân sự vẫn cao hơn Vertex hoặc Cloudflare ở cùng SLA.", base.BULLET_NUM_ID)

    # Trang 7: Lộ trình và nguồn.
    base.add_page_break(doc)
    base.add_heading(doc, "8. Lộ trình đề xuất", 1)
    base.add_table(
        doc,
        ["Giai đoạn", "Quy mô", "Hành động", "Cổng quyết định"],
        [
            ["Production hiện tại", "100-1.000", "Cloudflare primary; ghi nhận phút, p95, lỗi và chi phí", "Ổn định và có số liệu"],
            ["Benchmark", "500-1.000", "Golden set 300-500 audio; so Cloudflare, Google STT và custom model", "Chất lượng hơn >=10%"],
            ["Vertex shadow", "1.000-3.000", "Custom container + T4; thử warm và scale-to-zero", "TCO hơn >=20%"],
            ["Canary", ">=3.000", "Chuyển 5%-20%-50%; Cloudflare làm fallback", "30 ngày ổn định"],
            ["HA primary", ">=5.900", "Cân nhắc 2 replica Vertex hoặc self-host HA", "Đạt SLA và tải 2x"],
        ],
        [1550, 1500, 3900, 2410],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.7,
        header_fill=base.PALE_GREEN,
    )

    base.add_heading(doc, "9. Kết luận", 1)
    base.add_callout(
        doc,
        "QUYẾT ĐỊNH ĐỀ XUẤT",
        "Với 100, 500 và 1.000 người dùng, chọn Cloudflare. Vertex AI là bước trung gian hợp lý hơn self-host nếu cần custom model nhưng vẫn muốn hạ tầng, IAM và MLOps được quản lý. Google Speech-to-Text chỉ nên chọn khi chất lượng hoặc governance của Google tạo ra giá trị đủ lớn để bù chi phí theo phút.",
        fill=base.PALE_GREEN,
        accent=base.GREEN,
    )
    base.add_body(doc, "Đề xuất kiến trúc dài hạn: Cloudflare primary + Vertex shadow cho model tùy biến; chỉ chuyển lưu lượng khi vượt đủ bốn cổng chi phí, chất lượng, độ trễ và ổn định.")

    base.add_heading(doc, "10. Nguồn tham khảo", 1)
    base.add_source(doc, "Cloudflare Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/", "gói Workers Paid")
    base.add_source(doc, "Cloudflare Workers AI pricing", "https://developers.cloudflare.com/workers-ai/platform/pricing/", "quota và phí AI")
    base.add_source(doc, "Whisper Large V3 Turbo pricing", "https://developers.cloudflare.com/workers-ai/models/whisper-large-v3-turbo/", "0,00051 USD/phút")
    base.add_source(doc, "Google Vertex/Agent Platform pricing", "https://cloud.google.com/products/gemini-enterprise-agent-platform/pricing", "giá machine và accelerator theo vùng")
    base.add_source(doc, "Vertex deploy-model flags", "https://docs.cloud.google.com/sdk/gcloud/reference/beta/ai/endpoints/deploy-model", "autoscaling và scale-to-zero")
    base.add_source(doc, "Vertex custom container", "https://docs.cloud.google.com/vertex-ai/docs/predictions/use-custom-container", "đóng gói model inference")
    base.add_source(doc, "Vertex AI SLA", "https://cloud.google.com/vertex-ai/sla", "SLO cho custom prediction")
    base.add_source(doc, "Google Speech-to-Text pricing", "https://cloud.google.com/speech-to-text/pricing", "Standard và Dynamic Batch")
    base.add_source(doc, "Lambda GPU instances", "https://lambda.ai/instances", "giá GPU self-host tham khảo")
    base.add_source(doc, "SYSTRAN faster-whisper", "https://github.com/SYSTRAN/faster-whisper", "triển khai và benchmark")

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_report()

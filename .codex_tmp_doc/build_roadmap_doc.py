from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


OUTPUT = Path(__file__).resolve().parents[1] / "artifacts" / "lo_trinh_trien_khai_qwen_viet_anh.docx"

FONT = "Calibri"
INK = "1F2937"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
PALE = "F4F6F9"
WHITE = "FFFFFF"
GRID = "B8C4D1"


def set_run_font(run, *, name=FONT, size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Twips(width_dxa)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=120):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_paragraph_shading(paragraph, fill=PALE, left_border=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "20")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), left_border)
    p_bdr.append(left)


def create_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]

    def add_abstract(num_fmt, text, font=None):
        abstract_id = max(existing_abs + [-1]) + 1
        existing_abs.append(abstract_id)
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        lvl.append(fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "271")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        return abstract_id

    def add_num(abstract_id):
        num_id = max(existing_num + [0]) + 1
        existing_num.append(num_id)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        lvl_override = OxmlElement("w:lvlOverride")
        lvl_override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        lvl_override.append(start_override)
        num.append(lvl_override)
        numbering.append(num)
        return num_id

    bullet_abs = add_abstract("bullet", "•", FONT)
    decimal_abs = add_abstract("decimal", "%1.", FONT)
    return add_num(bullet_abs), (
        add_num(decimal_abs),
        add_num(decimal_abs),
        add_num(decimal_abs),
    )


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def add_bullet(doc, num_id, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    apply_numbering(p, num_id)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        set_run_font(p.add_run(text))
    return p


def add_numbered(doc, num_id, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    apply_numbering(p, num_id)
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_output(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_together = True
    set_run_font(p.add_run("Kết quả đầu ra: "), bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(text), color=INK)
    return p


def add_labeled_paragraph(doc, label, text):
    p = doc.add_paragraph(style="Normal")
    set_run_font(p.add_run(label + ": "), bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(text))
    return p


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.0
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = FONT
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Kicker" not in styles:
        kicker = styles.add_style("Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        kicker = styles["Kicker"]
    kicker.font.name = FONT
    kicker._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    kicker._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    kicker._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    kicker.font.size = Pt(10.5)
    kicker.font.bold = True
    kicker.font.color.rgb = RGBColor.from_string(BLUE)
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(6)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = False

    # Named layout override: footer-only page furniture for consistent rendering.
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.space_after = Pt(0)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    set_run_font(fp.add_run("Trang "), size=9, color=MUTED)
    add_page_field(fp)


def add_overview_table(doc):
    doc.add_heading("Tổng quan lộ trình", level=1)
    table = doc.add_table(rows=1, cols=3)
    headers = ("Giai đoạn", "Trọng tâm", "Đầu ra chính")
    widths = [1440, 3960, 3960]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=10.5, bold=True, color=NAVY)
    set_repeat_table_header(table.rows[0])

    rows = [
        ("1", "Đánh giá model, VRAM và tổng chi phí sở hữu", "Danh sách model phù hợp và phương án hạ tầng"),
        ("2", "Xây nền tảng kỹ năng vận hành", "Bộ bài thực hành Linux, Docker và GPU container"),
        ("3", "Xây kho dữ liệu Việt - Anh có kiểm duyệt", "Dataset phiên bản hóa, hướng dẫn gán nhãn và bộ đánh giá chuẩn"),
        ("4", "Self-host Qwen và benchmark trước khi huấn luyện", "Model server qua API cùng báo cáo baseline"),
        ("5", "Fine-tune bằng SFT kết hợp QLoRA", "Adapter, cấu hình huấn luyện và báo cáo đánh giá"),
        ("6", "Tích hợp Next.js và triển khai production", "Provider, fallback, monitoring và quy trình phát hành"),
    ]
    for stage, focus, output in rows:
        cells = table.add_row().cells
        values = (stage, focus, output)
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=10.2, color=INK, bold=(idx == 0))
    set_table_geometry(table, widths, indent=120)
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(0)


def build_document():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    configure_styles(doc)
    configure_section(doc.sections[0])
    bullet_num, decimal_nums = create_numbering(doc)
    data_decimal_num, training_decimal_num, fallback_decimal_num = decimal_nums

    props = doc.core_properties
    props.title = "Lộ trình xây dựng và triển khai trợ lý nói Việt - Anh"
    props.subject = "Qwen self-host, dữ liệu, fine-tune và tích hợp Next.js"
    props.author = "Nhóm dự án"
    props.keywords = "Qwen, self-host, QLoRA, SFT, Next.js, dữ liệu Việt Anh"

    kicker = doc.add_paragraph(style="Kicker")
    kicker.paragraph_format.space_before = Pt(18)
    set_run_font(kicker.add_run("ĐỀ CƯƠNG TRIỂN KHAI"), size=10.5, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("Lộ trình xây dựng và triển khai trợ lý nói Việt - Anh"), size=28, color=NAVY, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    set_run_font(
        subtitle.add_run("Từ lựa chọn mô hình, chuẩn bị dữ liệu và self-host Qwen đến fine-tune, tích hợp Next.js và vận hành production."),
        size=13,
        color=MUTED,
    )

    meta = doc.add_paragraph(style="Normal")
    meta.paragraph_format.space_after = Pt(2)
    set_run_font(meta.add_run("Đối tượng: "), size=10.5, bold=True, color=DARK_BLUE)
    set_run_font(meta.add_run("Nhóm kỹ thuật và quản lý sản phẩm  |  "), size=10.5, color=MUTED)
    set_run_font(meta.add_run("Phạm vi: "), size=10.5, bold=True, color=DARK_BLUE)
    set_run_font(meta.add_run("Pilot đến production"), size=10.5, color=MUTED)
    meta2 = doc.add_paragraph(style="Normal")
    meta2.paragraph_format.space_after = Pt(18)
    set_run_font(meta2.add_run("Phiên bản: "), size=10.5, bold=True, color=DARK_BLUE)
    set_run_font(meta2.add_run("1.0  |  Cập nhật: 20/08/2026"), size=10.5, color=MUTED)

    lead = doc.add_paragraph(style="Normal")
    lead.paragraph_format.left_indent = Inches(0.12)
    lead.paragraph_format.right_indent = Inches(0.08)
    lead.paragraph_format.space_before = Pt(2)
    lead.paragraph_format.space_after = Pt(14)
    lead.paragraph_format.keep_together = True
    set_paragraph_shading(lead)
    set_run_font(lead.add_run("Mục tiêu tổng thể. "), bold=True, color=NAVY)
    set_run_font(
        lead.add_run(
            "Xây dựng một nền tảng AI Việt - Anh có thể tự vận hành: dữ liệu được kiểm soát, mô hình được đo lường trước và sau fine-tune, backend có khả năng chuyển đổi nhà cung cấp, và hệ thống production có cơ chế giám sát, dự phòng và kiểm soát chi phí."
        ),
        color=INK,
    )

    add_overview_table(doc)

    doc.add_heading("Giai đoạn 1: Tìm hiểu mô hình, tham số và chi phí hạ tầng", level=1)
    add_labeled_paragraph(
        doc,
        "Mục tiêu",
        "Chọn được mô hình và cấu hình phục vụ phù hợp với chất lượng kỳ vọng, tải thực tế và ngân sách của dự án.",
    )
    doc.add_heading("Kiến thức về mô hình", level=2)
    for text, lead_text in [
        ("Kiến trúc và quy mô: hiểu số lượng tham số, độ dài ngữ cảnh, cơ chế attention, khả năng hỗ trợ tiếng Việt - tiếng Anh và giới hạn giấy phép sử dụng.", "Kiến trúc và quy mô:"),
        ("Độ chính xác số: so sánh BF16/FP16, INT8 và INT4; đánh giá ảnh hưởng của lượng tử hóa đến VRAM, tốc độ và chất lượng đầu ra.", "Độ chính xác số:"),
        ("Đặc tính phục vụ: đo thời gian phản hồi token đầu tiên, tốc độ sinh token, số yêu cầu đồng thời, độ ổn định và mức sử dụng GPU.", "Đặc tính phục vụ:"),
        ("Khả năng tinh chỉnh: kiểm tra chat template, tokenizer, hỗ trợ LoRA/QLoRA, yêu cầu dữ liệu và mức độ tương thích với công cụ huấn luyện.", "Khả năng tinh chỉnh:"),
    ]:
        add_bullet(doc, bullet_num, text, lead_text)

    doc.add_heading("Cách ước lượng VRAM", level=2)
    p = doc.add_paragraph(style="Normal")
    set_run_font(p.add_run("Ước lượng sơ bộ cho suy luận: "), bold=True, color=DARK_BLUE)
    set_run_font(p.add_run("VRAM cho trọng số ≈ số tham số × số bit mỗi tham số / 8. Sau đó cộng thêm KV cache, bộ đệm của runtime và phần dự phòng vận hành."))
    for text, lead_text in [
        ("Khi suy luận: KV cache tăng theo số phiên đồng thời, độ dài ngữ cảnh và kích thước batch; vì vậy phải benchmark trên tải gần với thực tế.", "Khi suy luận:"),
        ("Khi fine-tune: cần tính thêm activation, tham số adapter, gradient, trạng thái optimizer và chi phí do batch size, sequence length, gradient checkpointing tạo ra.", "Khi fine-tune:"),
        ("Nguyên tắc an toàn: không cấu hình sát trần VRAM; duy trì khoảng trống để tránh lỗi hết bộ nhớ khi lưu lượng hoặc độ dài đầu vào tăng đột biến.", "Nguyên tắc an toàn:"),
    ]:
        add_bullet(doc, bullet_num, text, lead_text)

    doc.add_heading("Các khoản chi phí phải tính", level=2)
    cost_items = [
        "GPU mua mới, khấu hao hoặc chi phí thuê theo giờ.",
        "Máy chủ CPU, RAM và SSD phục vụ model server, hàng đợi và các dịch vụ phụ trợ.",
        "Dung lượng tải model, lưu checkpoint, adapter và phiên bản dữ liệu.",
        "Điện năng, làm mát, không gian đặt máy hoặc phụ phí của nhà cung cấp hạ tầng.",
        "Băng thông mạng giữa backend, model server, kho lưu trữ và người dùng.",
        "Monitoring, tracing, cảnh báo, lưu log và thời gian lưu giữ dữ liệu quan sát.",
        "Backup dataset, checkpoint và adapter; chi phí kiểm tra khả năng khôi phục.",
        "Thời gian của kỹ sư cho triển khai, trực vận hành, xử lý sự cố và tối ưu hiệu năng.",
        "GPU hoặc phương án dự phòng khi máy chính lỗi, bảo trì hoặc quá tải.",
    ]
    last_cost = None
    for item in cost_items:
        last_cost = add_bullet(doc, bullet_num, item)
    if last_cost is not None:
        set_keep_with_next(last_cost)
    add_output(doc, "Bảng so sánh các model ứng viên, công thức ước lượng VRAM và tối thiểu ba kịch bản chi phí: pilot, tải mục tiêu và phương án dự phòng.")

    doc.add_heading("Giai đoạn 2: Xác định kiến thức cần học và nội dung có thể hoãn", level=1)
    add_labeled_paragraph(
        doc,
        "Mục tiêu",
        "Trang bị đủ năng lực để dựng, chạy, quan sát và xử lý lỗi của một model server GPU mà không sa vào các chủ đề chưa cần cho giai đoạn pilot.",
    )
    doc.add_heading("Nội dung cần học trước", level=2)
    required = [
        ("Linux: ", "quản lý file và quyền truy cập, tiến trình, service, biến môi trường, SSH, log, dung lượng đĩa và các lệnh chẩn đoán cơ bản."),
        ("Docker: ", "image, container, volume, network, Dockerfile, Compose, health check, log và chiến lược gắn phiên bản image."),
        ("Chạy model Hugging Face bằng Python: ", "tải model/tokenizer, áp dụng chat template, chọn dtype/device, sinh văn bản và đo thời gian, bộ nhớ."),
        ("Dựng Docker container có GPU: ", "NVIDIA driver và container runtime, ánh xạ GPU, kiểm tra bằng nvidia-smi, mount cache model và giới hạn tài nguyên."),
    ]
    for lead_text, tail in required:
        add_bullet(doc, bullet_num, lead_text + tail, lead_text)

    doc.add_heading("Nội dung có thể bỏ qua ở giai đoạn đầu", level=2)
    deferred = [
        "Tự huấn luyện foundation model từ đầu hoặc tự thiết kế kiến trúc transformer.",
        "Lập trình CUDA kernel và tối ưu mức phần cứng chuyên sâu.",
        "Huấn luyện phân tán nhiều node khi pilot vẫn chạy được trên một máy hoặc một cụm nhỏ.",
        "Kubernetes nâng cao trước khi lưu lượng và yêu cầu tự động mở rộng thực sự xuất hiện.",
        "Xây toàn bộ nền tảng MLOps tùy biến; ưu tiên công cụ tối giản nhưng có thể tái lập và quan sát được.",
    ]
    for item in deferred:
        add_bullet(doc, bullet_num, item)
    add_output(doc, "Bốn bài thực hành có thể chạy lại từ đầu, kèm runbook cài đặt, kiểm tra GPU, khởi động container và xử lý các lỗi thường gặp.")

    doc.add_heading("Giai đoạn 3: Xây dựng kho dữ liệu Việt - Anh có kiểm duyệt", level=1)
    add_labeled_paragraph(
        doc,
        "Mục tiêu",
        "Tạo một nguồn dữ liệu có nguồn gốc rõ ràng, phù hợp mục đích, an toàn với trẻ em và đủ nhất quán để dùng cho SFT cũng như đánh giá.",
    )
    doc.add_heading("Hướng dẫn gán nhãn", level=2)
    guide_items = [
        "Mục tiêu của từng nhãn, định nghĩa trường dữ liệu và tiêu chí chấp nhận hoặc loại bỏ mẫu.",
        "Quy tắc về độ chính xác ngôn ngữ, tính tự nhiên, mức độ phù hợp lứa tuổi, tính hữu ích và an toàn.",
        "Ví dụ đúng, ví dụ sai và các tình huống biên để giảm khác biệt giữa người gán nhãn.",
        "Quy trình xử lý bất đồng, kiểm tra chéo, đo mức nhất quán và cập nhật phiên bản hướng dẫn.",
        "Quy tắc không đưa thông tin nhận dạng cá nhân, nội dung nhạy cảm hoặc dữ liệu không đủ quyền sử dụng vào tập huấn luyện.",
    ]
    for item in guide_items:
        add_bullet(doc, bullet_num, item)

    doc.add_heading("Yêu cầu riêng đối với dữ liệu trẻ em", level=2)
    child_items = [
        ("Quyền và sự đồng ý: ", "chỉ sử dụng dữ liệu có căn cứ pháp lý, quyền khai thác và cơ chế đồng ý phù hợp; lưu bằng chứng về nguồn gốc dữ liệu."),
        ("Bảo vệ danh tính: ", "loại bỏ tên thật, địa chỉ, trường học, số điện thoại, giọng nói nhận dạng được và các dấu hiệu định danh khác khi không thật sự cần thiết."),
        ("Phù hợp lứa tuổi: ", "kiểm soát từ vựng, chủ đề, mức độ khó, cách phản hồi và cơ chế từ chối đối với nội dung nguy hiểm hoặc không phù hợp."),
        ("Tính đại diện: ", "cân bằng độ tuổi, vùng miền, giới tính, mức độ thành thạo và các kiểu phát âm; tránh củng cố định kiến."),
        ("Kiểm soát truy cập: ", "phân quyền kho dữ liệu, mã hóa khi lưu/truyền, ghi nhận truy cập, quy định thời gian lưu và quy trình xóa dữ liệu."),
    ]
    for lead_text, tail in child_items:
        add_bullet(doc, bullet_num, lead_text + tail, lead_text)

    doc.add_heading("Quy trình dữ liệu đề xuất", level=2)
    for text, lead_text in [
        ("Thu thập và kiểm tra quyền sử dụng dữ liệu.", "Thu thập"),
        ("Lọc trùng, chuẩn hóa định dạng và loại dữ liệu hỏng hoặc ngoài phạm vi.", "Lọc trùng"),
        ("Ẩn danh, quét thông tin cá nhân và đánh dấu nội dung nhạy cảm.", "Ẩn danh"),
        ("Gán nhãn, kiểm tra chéo và xử lý bất đồng theo hướng dẫn.", "Gán nhãn"),
        ("Chia train/validation/test theo nguyên tắc tránh rò rỉ nội dung hoặc người nói.", "Chia train"),
        ("Phiên bản hóa dataset, dataset card, thống kê chất lượng và lịch sử thay đổi.", "Phiên bản hóa"),
    ]:
        add_numbered(doc, data_decimal_num, text)
    add_output(doc, "Hướng dẫn gán nhãn, dataset card, tập dữ liệu đã phiên bản hóa, báo cáo kiểm duyệt và một bộ đánh giá chuẩn được khóa phiên bản.")

    doc.add_heading("Giai đoạn 4: Self-host Qwen và benchmark trước khi train", level=1)
    add_labeled_paragraph(
        doc,
        "Mục tiêu",
        "Tạo baseline có thể đo lường để xác định giới hạn của model gốc và chỉ fine-tune khi đã biết rõ vấn đề cần cải thiện.",
    )
    doc.add_heading("Kiến trúc pilot", level=2)
    architecture = doc.add_paragraph(style="Normal")
    architecture.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    architecture.paragraph_format.space_before = Pt(4)
    architecture.paragraph_format.space_after = Pt(10)
    architecture.paragraph_format.keep_together = True
    set_paragraph_shading(architecture, fill=LIGHT_BLUE, left_border=DARK_BLUE)
    set_run_font(
        architecture.add_run("Ứng dụng/Backend Next.js  ->  Provider nội bộ  ->  Qwen Model Server  ->  GPU\nMonitoring, log và benchmark thu thập tín hiệu xuyên suốt luồng xử lý."),
        size=10.5,
        color=NAVY,
        bold=True,
    )

    doc.add_heading("Dựng model server", level=2)
    server_items = [
        "Đóng gói inference server trong container, gắn phiên bản rõ ràng cho model, runtime và image.",
        "Cung cấp API hội thoại ổn định, ưu tiên giao diện tương thích OpenAI để đơn giản hóa tích hợp provider.",
        "Bổ sung health/readiness endpoint, xác thực nội bộ, giới hạn kích thước đầu vào, timeout và quản lý số phiên đồng thời.",
        "Tách model cache và log khỏi filesystem tạm của container; bảo đảm có quy trình khởi động lại và nạp model có kiểm soát.",
    ]
    for item in server_items:
        add_bullet(doc, bullet_num, item)

    doc.add_heading("Benchmark baseline", level=2)
    benchmark_items = [
        ("Hiệu năng: ", "thời gian phản hồi token đầu tiên, tổng thời gian trả lời, token/giây, throughput, mức dùng VRAM và tỷ lệ lỗi."),
        ("Chất lượng: ", "độ tự nhiên tiếng Việt - tiếng Anh, độ chính xác theo nhiệm vụ, khả năng tuân thủ hướng dẫn và xử lý hội thoại nhiều lượt."),
        ("An toàn: ", "phản hồi với dữ liệu cá nhân, nội dung trẻ em, yêu cầu nguy hiểm và tình huống model cần từ chối hoặc chuyển hướng."),
        ("Độ bền: ", "tải đồng thời, prompt dài, kết nối bị ngắt, model restart và hành vi khi GPU gần hết bộ nhớ."),
    ]
    for lead_text, tail in benchmark_items:
        add_bullet(doc, bullet_num, lead_text + tail, lead_text)
    add_output(doc, "Model server gọi được qua API, bộ script benchmark có thể chạy lặp lại, báo cáo baseline và danh sách khoảng trống chất lượng cần giải quyết.")

    doc.add_heading("Giai đoạn 5: Fine-tune bằng SFT kết hợp QLoRA", level=1)
    add_labeled_paragraph(
        doc,
        "Mục tiêu",
        "Tinh chỉnh hành vi và phong cách hội thoại trên dữ liệu mục tiêu với chi phí VRAM thấp hơn, đồng thời giữ model gốc và adapter tách biệt để dễ thử nghiệm và hoàn tác.",
    )
    doc.add_heading("Cơ chế fine-tune", level=2)
    mechanisms = [
        ("SFT: ", "dạy model bắt chước câu trả lời chuẩn từ các cặp hội thoại đã kiểm duyệt; chất lượng dữ liệu và cách tạo nhãn loss ảnh hưởng trực tiếp đến kết quả."),
        ("QLoRA: ", "lượng tử hóa model nền trong quá trình huấn luyện và chỉ cập nhật các ma trận adapter LoRA, từ đó giảm đáng kể bộ nhớ cần dùng so với fine-tune toàn bộ tham số."),
        ("Nguyên tắc triển khai: ", "giữ nguyên checkpoint nền, lưu adapter theo phiên bản, ghi lại seed và cấu hình để có thể tái lập hoặc rollback."),
    ]
    for lead_text, tail in mechanisms:
        add_bullet(doc, bullet_num, lead_text + tail, lead_text)

    doc.add_heading("Định dạng hội thoại", level=2)
    add_bullet(doc, bullet_num, "Chuẩn hóa mỗi mẫu theo vai trò system, user và assistant; sử dụng đúng chat template/tokenizer của model Qwen được chọn.")
    add_bullet(doc, bullet_num, "Quy định rõ điểm kết thúc lượt, xử lý hội thoại nhiều lượt, độ dài tối đa và cách loại các mẫu thiếu ngữ cảnh hoặc mâu thuẫn.")
    add_bullet(doc, bullet_num, "Chỉ tính loss trên phần câu trả lời của assistant khi thiết kế huấn luyện yêu cầu như vậy; kiểm thử dữ liệu sau tokenize để tránh nhãn lệch hoặc mất token kết thúc.")
    add_bullet(doc, bullet_num, "Giữ một bộ test cố định, không xuất hiện trong train, để so sánh model gốc và các adapter theo cùng một tiêu chí.")

    example = doc.add_paragraph(style="Normal")
    example.paragraph_format.left_indent = Inches(0.18)
    example.paragraph_format.right_indent = Inches(0.12)
    example.paragraph_format.space_before = Pt(4)
    example.paragraph_format.space_after = Pt(10)
    example.paragraph_format.keep_together = True
    set_paragraph_shading(example, fill=PALE, left_border=BLUE)
    set_run_font(example.add_run("Ví dụ cấu trúc logic\n"), size=10.5, color=NAVY, bold=True)
    set_run_font(example.add_run("system: Quy tắc vai trò và an toàn\nuser: Câu hỏi hoặc lời nói của người học\nassistant: Phản hồi Việt - Anh đã được kiểm duyệt"), size=10, color=INK)

    doc.add_heading("Quy trình huấn luyện và đánh giá", level=2)
    train_steps = [
        "Khóa phiên bản dataset, chat template, model nền và bộ đánh giá.",
        "Chạy thử trên tập nhỏ để kiểm tra loss, tokenize, checkpoint và khả năng resume.",
        "Thực hiện SFT + QLoRA với cấu hình được ghi nhận đầy đủ; theo dõi train/validation loss và tài nguyên GPU.",
        "So sánh model gốc với adapter trên benchmark hiệu năng, chất lượng và an toàn.",
        "Chỉ chọn adapter khi cải thiện mục tiêu chính mà không làm suy giảm đáng kể các tiêu chí quan trọng khác.",
    ]
    for step in train_steps:
        add_numbered(doc, training_decimal_num, step)
    add_output(doc, "Adapter đã phiên bản hóa, cấu hình và script huấn luyện tái lập được, log/checkpoint, model card và báo cáo so sánh trước - sau fine-tune.")

    doc.add_heading("Giai đoạn 6: Tích hợp backend Next.js và triển khai production", level=1)
    add_labeled_paragraph(
        doc,
        "Mục tiêu",
        "Tách biệt ứng dụng khỏi chi tiết của model server, bảo vệ cấu hình nhạy cảm và duy trì dịch vụ khi model chính chậm, lỗi hoặc quá tải.",
    )
    doc.add_heading("Thiết kế provider", level=2)
    provider_items = [
        "Định nghĩa interface chung cho request, response, streaming, usage, timeout và mã lỗi; ứng dụng không gọi trực tiếp model server.",
        "Cài đặt provider cho Qwen self-host và provider dự phòng; chuẩn hóa lịch sử hội thoại, system prompt và giới hạn token tại một nơi.",
        "Tách logic retry, circuit breaker, rate limit và quan sát telemetry khỏi logic nghiệp vụ của API route.",
    ]
    for item in provider_items:
        add_bullet(doc, bullet_num, item)

    doc.add_heading("Cấu hình server-side", level=2)
    config_items = [
        "Lưu endpoint, API key, model ID, timeout, ngưỡng retry và cờ bật/tắt fallback trong biến môi trường chỉ dành cho server.",
        "Không đưa bí mật hoặc địa chỉ nội bộ vào client bundle; kiểm tra cấu hình ngay khi ứng dụng khởi động và thất bại sớm nếu thiếu giá trị bắt buộc.",
        "Phân tách cấu hình development, staging và production; quản lý secret bằng dịch vụ chuyên dụng và xoay vòng định kỳ.",
    ]
    for item in config_items:
        add_bullet(doc, bullet_num, item)

    doc.add_heading("Luồng fallback", level=2)
    fallback_steps = [
        "Gửi yêu cầu tới Qwen self-host với timeout và mã truy vết duy nhất.",
        "Nếu lỗi tạm thời, retry có backoff trong giới hạn cho phép; không retry các lỗi dữ liệu hoặc xác thực.",
        "Khi vượt ngưỡng lỗi hoặc độ trễ, mở circuit breaker và chuyển sang provider dự phòng theo chính sách.",
        "Nếu cả hai provider không khả dụng, trả thông báo an toàn, không làm mất trạng thái hội thoại và ghi sự kiện để điều tra.",
        "Khi model chính phục hồi, thực hiện health probe và đưa lưu lượng trở lại từng bước.",
    ]
    for step in fallback_steps:
        add_numbered(doc, fallback_decimal_num, step)

    doc.add_heading("Triển khai production", level=2)
    production_items = [
        ("Tách dịch vụ: ", "backend Next.js và model server GPU triển khai độc lập để có thể mở rộng, cập nhật và rollback riêng."),
        ("Mạng và bảo mật: ", "TLS, mạng nội bộ, xác thực giữa dịch vụ, giới hạn truy cập, quét image và nguyên tắc đặc quyền tối thiểu."),
        ("Quan sát hệ thống: ", "metrics, log có cấu trúc, tracing, dashboard và cảnh báo cho độ trễ, tỷ lệ lỗi, hàng đợi, VRAM/GPU và chi phí."),
        ("Độ tin cậy: ", "health/readiness probe, graceful shutdown, rolling deployment, backup cấu hình/adapter và diễn tập khôi phục."),
        ("Kiểm soát phát hành: ", "staging, canary hoặc A/B test, bộ smoke test tự động và tiêu chí rollback rõ ràng."),
        ("Quản trị chi phí: ", "theo dõi chi phí theo môi trường, tải và số token; đặt cảnh báo ngân sách và lịch tắt tài nguyên không cần thiết."),
    ]
    for lead_text, tail in production_items:
        add_bullet(doc, bullet_num, lead_text + tail, lead_text)
    add_output(doc, "Provider hoàn chỉnh, cấu hình server-side an toàn, luồng fallback đã kiểm thử, dashboard/cảnh báo, runbook vận hành và quy trình triển khai - rollback.")

    doc.add_heading("Tiêu chí hoàn thành toàn bộ lộ trình", level=1)
    completion = [
        "Model và hạ tầng được chọn dựa trên benchmark và tổng chi phí sở hữu, không chỉ dựa trên số lượng tham số.",
        "Dữ liệu có nguồn gốc, quyền sử dụng, kiểm duyệt, hướng dẫn gán nhãn và phiên bản rõ ràng; dữ liệu trẻ em được bảo vệ theo chính sách riêng.",
        "Qwen self-host phục vụ ổn định qua API và có baseline hiệu năng, chất lượng, an toàn trước khi fine-tune.",
        "Adapter SFT + QLoRA cải thiện mục tiêu đã xác định và có thể tái lập, so sánh, triển khai hoặc rollback độc lập.",
        "Backend Next.js không phụ thuộc cứng vào một model; fallback, monitoring, backup và quy trình xử lý sự cố đã được kiểm thử.",
        "Hệ thống production có chủ sở hữu, SLO, cảnh báo, ngân sách và lịch rà soát định kỳ.",
    ]
    for item in completion:
        add_bullet(doc, bullet_num, item)

    closing = doc.add_paragraph(style="Normal")
    closing.paragraph_format.space_before = Pt(8)
    closing.paragraph_format.space_after = Pt(0)
    closing.paragraph_format.keep_together = True
    set_paragraph_shading(closing, fill=LIGHT_BLUE, left_border=NAVY)
    set_run_font(closing.add_run("Nguyên tắc xuyên suốt. "), bold=True, color=NAVY)
    set_run_font(closing.add_run("Đo trước khi tối ưu, kiểm soát dữ liệu trước khi huấn luyện, và thiết kế khả năng phục hồi trước khi đưa hệ thống vào production."), color=INK)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())

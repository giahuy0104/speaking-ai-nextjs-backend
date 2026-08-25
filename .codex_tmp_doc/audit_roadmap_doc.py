import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOCX = Path(__file__).resolve().parents[1] / "artifacts" / "lo_trinh_trien_khai_qwen_viet_anh.docx"


def close(a, b, tolerance=1000):
    return abs(a - b) <= tolerance


doc = Document(DOCX)
assert len(doc.sections) == 1
section = doc.sections[0]
assert close(section.page_width, Inches(8.5))
assert close(section.page_height, Inches(11))
assert all(
    close(value, Inches(1))
    for value in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin)
)
assert close(section.header_distance, Inches(0.492))
assert close(section.footer_distance, Inches(0.492))
assert not section.header.paragraphs[0].text.strip()
assert section.footer.paragraphs[0].text.startswith("Trang ")

styles = doc.styles
normal = styles["Normal"]
assert normal.font.name == "Calibri"
assert close(normal.font.size, Pt(11))
assert close(normal.paragraph_format.space_before or 0, Pt(0))
assert close(normal.paragraph_format.space_after, Pt(6))
assert abs(float(normal.paragraph_format.line_spacing) - 1.25) < 0.001

expected_headings = {
    "Heading 1": (16, "2E74B5", 18, 10),
    "Heading 2": (13, "2E74B5", 14, 7),
    "Heading 3": (12, "1F4D78", 10, 5),
}
for name, (size, color, before, after) in expected_headings.items():
    style = styles[name]
    assert style.font.name == "Calibri"
    assert close(style.font.size, Pt(size))
    assert str(style.font.color.rgb) == color
    assert close(style.paragraph_format.space_before, Pt(before))
    assert close(style.paragraph_format.space_after, Pt(after))
    assert style.paragraph_format.keep_with_next

assert len(doc.tables) == 1
table = doc.tables[0]
assert len(table.rows) == 7 and len(table.columns) == 3
tbl_pr = table._tbl.tblPr
assert tbl_pr.find(qn("w:tblW")).get(qn("w:w")) == "9360"
assert tbl_pr.find(qn("w:tblW")).get(qn("w:type")) == "dxa"
assert tbl_pr.find(qn("w:tblInd")).get(qn("w:w")) == "120"
assert tbl_pr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
grid = [int(e.get(qn("w:w"))) for e in table._tbl.tblGrid.findall(qn("w:gridCol"))]
assert grid == [1440, 3960, 3960]
for row in table.rows:
    for idx, cell in enumerate(row.cells):
        tc_pr = cell._tc.get_or_add_tcPr()
        assert int(tc_pr.find(qn("w:tcW")).get(qn("w:w"))) == grid[idx]
        tc_mar = tc_pr.find(qn("w:tcMar"))
        margins = {name: int(tc_mar.find(qn(f"w:{name}")).get(qn("w:w"))) for name in ("top", "bottom", "start", "end")}
        assert margins == {"top": 80, "bottom": 80, "start": 120, "end": 120}

numbered_paragraphs = []
all_text = []
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    all_text.append(text)
    num_pr = paragraph._p.get_or_add_pPr().find(qn("w:numPr"))
    if num_pr is not None:
        numbered_paragraphs.append(paragraph)
    assert not text.startswith(("- ", "• ", "● "))
    assert not re.match(r"^\d+\.\s", text)

assert len(numbered_paragraphs) >= 60
used_num_ids = {
    int(p._p.get_or_add_pPr().find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val")))
    for p in numbered_paragraphs
}
assert len(used_num_ids) == 4

numbering = doc.part.numbering_part.element
num_map = {
    int(num.get(qn("w:numId"))): int(num.find(qn("w:abstractNumId")).get(qn("w:val")))
    for num in numbering.findall(qn("w:num"))
}
abstracts = {
    int(abs_num.get(qn("w:abstractNumId"))): abs_num
    for abs_num in numbering.findall(qn("w:abstractNum"))
}
for num_id in used_num_ids:
    abstract = abstracts[num_map[num_id]]
    lvl = abstract.find(qn("w:lvl"))
    p_pr = lvl.find(qn("w:pPr"))
    ind = p_pr.find(qn("w:ind"))
    spacing = p_pr.find(qn("w:spacing"))
    assert ind.get(qn("w:left")) == "540"
    assert ind.get(qn("w:hanging")) == "271"
    assert spacing.get(qn("w:after")) == "80"
    assert spacing.get(qn("w:line")) == "300"
    assert spacing.get(qn("w:lineRule")) == "auto"

decimal_ids = [
    num_id
    for num_id in used_num_ids
    if abstracts[num_map[num_id]].find(qn("w:lvl")).find(qn("w:numFmt")).get(qn("w:val")) == "decimal"
]
assert len(decimal_ids) == 3
for num_id in decimal_ids:
    num = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == num_id)
    override = num.find(qn("w:lvlOverride"))
    assert override is not None
    assert override.find(qn("w:startOverride")).get(qn("w:val")) == "1"

joined = "\n".join(all_text)
for forbidden in (":codex-file-citation", "TODO", "{{", "}}", "PLACEHOLDER"):
    assert forbidden not in joined

with zipfile.ZipFile(DOCX) as zf:
    styles_xml = zf.read("word/styles.xml").decode("utf-8")
    title_block = re.search(r'<w:style[^>]*w:styleId="Title".*?</w:style>', styles_xml).group(0)
    assert "<w:pBdr" not in title_block

print(
    {
        "status": "PASS",
        "paragraphs": len(doc.paragraphs),
        "numbered_paragraphs": len(numbered_paragraphs),
        "used_num_ids": sorted(used_num_ids),
        "tables": len(doc.tables),
        "file_bytes": DOCX.stat().st_size,
    }
)

import json
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent):
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if child.tag.endswith("}p"):
            yield "paragraph", Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield "table", Table(child, parent)


def cell_text(cell):
    return "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip()).strip()


def main():
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(source)
    blocks = []
    table_number = 0
    paragraph_number = 0
    for kind, item in iter_blocks(doc):
        if kind == "paragraph":
            paragraph_number += 1
            text = item.text.strip()
            if text:
                blocks.append({
                    "type": "paragraph",
                    "number": paragraph_number,
                    "style": item.style.name if item.style else "",
                    "text": text,
                })
        else:
            table_number += 1
            rows = [[cell_text(cell) for cell in row.cells] for row in item.rows]
            blocks.append({
                "type": "table",
                "number": table_number,
                "rows": rows,
                "row_count": len(rows),
                "column_count": max((len(row) for row in rows), default=0),
            })

    sections = []
    for index, section in enumerate(doc.sections, start=1):
        sections.append({
            "number": index,
            "page_width": section.page_width,
            "page_height": section.page_height,
            "top_margin": section.top_margin,
            "bottom_margin": section.bottom_margin,
            "left_margin": section.left_margin,
            "right_margin": section.right_margin,
            "header": [p.text for p in section.header.paragraphs if p.text.strip()],
            "footer": [p.text for p in section.footer.paragraphs if p.text.strip()],
        })

    payload = {
        "source": str(source),
        "paragraph_count": paragraph_number,
        "table_count": table_number,
        "inline_shape_count": len(doc.inline_shapes),
        "sections": sections,
        "blocks": blocks,
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
